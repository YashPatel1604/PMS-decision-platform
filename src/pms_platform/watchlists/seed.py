"""Seed a watchlist from Research Fair Value Excel (company names only)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from pms_platform.models.watchlist import Watchlist
from pms_platform.research_paths import research_portfolio_dir
from pms_platform.watchlists import service as wl
from pms_platform.watchlists.service import DuplicateMemberError, WatchlistError

DEFAULT_WATCHLIST_NAME = "Fair Value"
RESEARCH_FILENAME = "Stocks_FairValue_Watchlist.xlsx"
_SKIP_LABELS = frozenset(
    {
        "company",
        "competitors",
        "previous owned",
        "total",
        "totals",
    }
)


def research_fair_value_watchlist_path() -> Path | None:
    """Return Research/Portfolio/Stocks_FairValue_Watchlist.xlsx when present."""
    portfolio = research_portfolio_dir()
    if portfolio is None:
        return None
    path = portfolio / RESEARCH_FILENAME
    return path if path.is_file() else None


def unique_company_names(path: Path) -> list[str]:
    """Union of Col A company names across sheets, first-seen order, junk skipped."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ordered: list[str] = []
        seen: set[str] = set()
        for sheet_name in workbook.sheetnames:
            for (value,) in workbook[sheet_name].iter_rows(
                min_col=1, max_col=1, values_only=True
            ):
                name = _clean_company_name(value)
                if name is None:
                    continue
                key = name.casefold()
                if key in seen:
                    continue
                seen.add(key)
                ordered.append(name)
        return ordered
    finally:
        workbook.close()


def _clean_company_name(value: object) -> str | None:
    if value is None or hasattr(value, "year"):
        return None
    if isinstance(value, int | float):
        return None
    text = str(value).strip()
    if not text or not any(ch.isalpha() for ch in text):
        return None
    if text.casefold() in _SKIP_LABELS:
        return None
    return text


@dataclass(frozen=True)
class SeedResult:
    watchlist_id: int
    watchlist_name: str
    added: int
    skipped: int
    names: int


def seed_watchlist(
    session: Session,
    *,
    path: Path,
    name: str | None = None,
    force: bool = False,
) -> SeedResult:
    """Add unique Excel names to the default (or named) watchlist."""
    if not path.is_file():
        raise WatchlistError(f"Watchlist Excel not found: {path}")
    names = unique_company_names(path)
    if not names:
        raise WatchlistError(f"No company names found in {path}")

    watchlist = _target_watchlist(session, name)
    existing = wl.member_count(session, watchlist.watchlist_id)
    if existing > 0 and not force:
        raise WatchlistError(
            f"Watchlist '{watchlist.name}' already has {existing} members; "
            "pass --force to add missing names"
        )

    added = 0
    skipped = 0
    for company in names:
        try:
            wl.add_member(
                session,
                watchlist.watchlist_id,
                wl.MemberInput(portfolio_name=company, display_name=company),
            )
            added += 1
        except DuplicateMemberError:
            skipped += 1
    session.flush()
    return SeedResult(
        watchlist_id=watchlist.watchlist_id,
        watchlist_name=watchlist.name,
        added=added,
        skipped=skipped,
        names=len(names),
    )


def _target_watchlist(session: Session, name: str | None) -> Watchlist:
    if name and name.strip():
        clean = name.strip()
        existing = session.scalar(select(Watchlist).where(Watchlist.name == clean))
        if existing is not None:
            return existing
        return wl.create_watchlist(
            session,
            name=clean,
            description="Seeded from Research Fair Value watchlist",
        )
    default = session.scalar(select(Watchlist).where(Watchlist.is_default.is_(True)))
    if default is not None:
        return default
    return wl.create_watchlist(
        session,
        name=DEFAULT_WATCHLIST_NAME,
        description="Seeded from Research Fair Value watchlist",
    )
