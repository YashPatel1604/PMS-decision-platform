"""Harvest Dad's Portfolio_*.xlsx Change notes into the transactions master.

Change column examples:
  Add 550@683.63 = 375997
  Buy 532@525.87 = 279763
  Sell 130@355.29 = 46188
  Sell 1840@169.66

Liquid / Split / free-text notes are skipped (reported), not invented into equity txs.
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from pms_platform.ingestion.common import is_liquid_holding
from pms_platform.ingestion.snapshots import SKIP_SHEETS, _sheet_snapshot_date
from pms_platform.masters.apply import _append_transaction_row, _sync_to_raw
from pms_platform.masters.paths import MasterKind, resolve_master_path
from pms_platform.models.enums import EventType
from pms_platform.research_paths import portfolio_snapshot_source_dirs, research_portfolio_dir

_TRADE_RE = re.compile(
    r"^\s*(?P<kind>add|buy|sell)\s+"
    r"(?P<qty>\d[\d,]*)\s*"
    r"(?:@\s*(?P<price>\d+(?:\.\d+)?))?\s*"
    r"(?:=\s*(?P<amount>\d[\d,]*(?:\.\d+)?))?\s*$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class HarvestedTrade:
    workbook_name: str
    sheet_name: str
    event_date: date
    portfolio_name: str
    event_type: str  # Buy / Sell
    quantity: int
    price: Decimal | None
    amount: Decimal | None
    raw_change: str
    source_note: str
    fingerprint: str


@dataclass
class HarvestResult:
    candidates: list[HarvestedTrade] = field(default_factory=list)
    already_present: list[HarvestedTrade] = field(default_factory=list)
    appended: list[HarvestedTrade] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    backup_path: str | None = None
    master_path: str | None = None
    synced_raw: str | None = None


def _dec(raw: str | None) -> Decimal | None:
    if raw is None or not str(raw).strip():
        return None
    try:
        return Decimal(str(raw).replace(",", "").strip())
    except InvalidOperation:
        return None


def parse_change_note(raw: str) -> tuple[str, int, Decimal | None, Decimal | None] | None:
    """Return (Buy|Sell, qty, price, amount) or None if not an equity trade note."""
    text = str(raw).strip()
    if not text:
        return None
    match = _TRADE_RE.match(text)
    if match is None:
        return None
    kind = match.group("kind").lower()
    qty = int(match.group("qty").replace(",", ""))
    if qty <= 0:
        return None
    price = _dec(match.group("price"))
    amount = _dec(match.group("amount"))
    if kind == "sell":
        event = EventType.SELL.value
    else:
        event = EventType.BUY.value
    # Equity buys/sells need a price when amount is present without @ — rare.
    if price is None and amount is not None and qty > 0:
        price = (amount / Decimal(qty)).quantize(Decimal("0.01"))
    if price is None:
        # Sell/Buy qty-only (Liquid-style) — not enough for equity master row.
        return None
    if amount is None:
        amount = (Decimal(qty) * price).quantize(Decimal("0.01"))
    return event, qty, price, amount


def _fingerprint(
    workbook_name: str,
    sheet_name: str,
    portfolio_name: str,
    event_type: str,
    quantity: int,
    price: Decimal | None,
) -> str:
    price_part = f"{price:f}" if price is not None else ""
    return (
        f"{workbook_name}|{sheet_name}|{portfolio_name}|{event_type}|"
        f"{quantity}|{price_part}"
    )


def iter_portfolio_workbooks(extra: Path | None = None) -> list[Path]:
    paths: list[Path] = []
    seen: set[str] = set()
    dirs = list(portfolio_snapshot_source_dirs())
    portfolio = research_portfolio_dir()
    if portfolio is not None and portfolio not in dirs:
        dirs.insert(0, portfolio)
    if extra is not None:
        dirs.insert(0, extra if extra.is_dir() else extra.parent)
        if extra.is_file():
            return [extra]
    for directory in dirs:
        for path in sorted(directory.glob("Portfolio_*.xlsx")):
            if path.name.startswith("~$"):
                continue
            if path.name in seen:
                continue
            seen.add(path.name)
            paths.append(path)
    return paths


def harvest_change_notes(workbooks: list[Path] | None = None) -> list[HarvestedTrade]:
    """Parse equity Buy/Sell candidates from Portfolio Change columns."""
    books = workbooks if workbooks is not None else iter_portfolio_workbooks()
    out: list[HarvestedTrade] = []
    for path in books:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet_name in workbook.sheetnames:
                if sheet_name.strip().lower() in SKIP_SHEETS:
                    continue
                worksheet = workbook[sheet_name]
                event_date = _sheet_snapshot_date(path, sheet_name, worksheet)
                if event_date is None:
                    continue
                for row in worksheet.iter_rows(min_row=5, values_only=True):
                    if not row or row[0] is None:
                        continue
                    stock = str(row[0]).strip()
                    if not stock or stock.lower() in {"total", "cash", "stock"}:
                        continue
                    if is_liquid_holding(stock):
                        continue
                    change = row[6] if len(row) > 6 else None
                    if change is None or not str(change).strip():
                        continue
                    raw = str(change).strip()
                    parsed = parse_change_note(raw)
                    if parsed is None:
                        continue
                    event_type, qty, price, amount = parsed
                    fp = _fingerprint(
                        path.name, sheet_name, stock, event_type, qty, price
                    )
                    note = f"{path.name} / {sheet_name}: {raw}"
                    out.append(
                        HarvestedTrade(
                            workbook_name=path.name,
                            sheet_name=sheet_name,
                            event_date=event_date,
                            portfolio_name=stock,
                            event_type=event_type,
                            quantity=qty,
                            price=price,
                            amount=amount,
                            raw_change=raw,
                            source_note=note,
                            fingerprint=fp,
                        )
                    )
        finally:
            workbook.close()
    return out


def _load_existing_keys(master_path: Path) -> tuple[set[str], set[tuple]]:
    """Fingerprints from Source/Notes + (name, date, type, qty, price) tuples."""
    fingerprints: set[str] = set()
    tuples: set[tuple] = set()
    workbook = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        current_stock: str | None = None
        for row in sheet.iter_rows(min_row=2, values_only=True):
            stock, event_date, event_label, quantity, price, _amount, note = (
                (row + (None,) * 8)[1:8]
            )
            if stock is not None and str(stock).strip():
                current_stock = str(stock).strip()
            if current_stock is None or event_date is None or event_label is None:
                continue
            if quantity is None:
                continue
            qty = int(quantity)
            price_d = (
                Decimal(str(price)) if price is not None and str(price).strip() else None
            )
            event = str(event_label).strip()
            day = event_date.date() if isinstance(event_date, datetime) else event_date
            if not isinstance(day, date):
                continue
            tuples.add((current_stock.lower(), day, event.lower(), qty, price_d))
            if note is not None and str(note).strip():
                text = str(note).strip()
                # Prefer explicit fingerprint segment if we stored workbook|sheet|…
                if "|" in text and text.count("|") >= 4:
                    fingerprints.add(text.split(";")[0].strip())
                # Also match "Portfolio_2026.xlsx / 16Apr26: Add …"
                m = re.search(
                    r"(Portfolio_\d{4}\.xlsx)\s*/\s*([^:]+):\s*(.+)$",
                    text,
                    re.IGNORECASE,
                )
                if m:
                    wb, sheet_n, change = m.group(1), m.group(2).strip(), m.group(3).strip()
                    parsed = parse_change_note(change)
                    if parsed is not None:
                        et, q, p, _a = parsed
                        fingerprints.add(
                            _fingerprint(wb, sheet_n, current_stock, et, q, p)
                        )
    finally:
        workbook.close()
    return fingerprints, tuples


def _is_present(
    trade: HarvestedTrade,
    fingerprints: set[str],
    tuples: set[tuple],
) -> bool:
    if trade.fingerprint in fingerprints:
        return True
    key = (
        trade.portfolio_name.lower(),
        trade.event_date,
        trade.event_type.lower(),
        trade.quantity,
        trade.price,
    )
    return key in tuples


def apply_harvest(
    *,
    dry_run: bool = True,
    workbooks: list[Path] | None = None,
    master_path: Path | None = None,
) -> HarvestResult:
    """Propose (and optionally append) missing equity trades from Change notes."""
    result = HarvestResult()
    path = master_path or resolve_master_path(MasterKind.TRANSACTIONS)
    result.master_path = str(path)
    if not path.is_file():
        result.skipped.append(f"Missing transactions master: {path}")
        return result

    candidates = harvest_change_notes(workbooks)
    result.candidates = candidates
    fingerprints, tuples = _load_existing_keys(path)

    missing: list[HarvestedTrade] = []
    for trade in candidates:
        if _is_present(trade, fingerprints, tuples):
            result.already_present.append(trade)
        else:
            missing.append(trade)

    # Unparsed change notes (for report)
    for book in workbooks if workbooks is not None else iter_portfolio_workbooks():
        wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
        try:
            for sheet_name in wb.sheetnames:
                if sheet_name.strip().lower() in SKIP_SHEETS:
                    continue
                ws = wb[sheet_name]
                if _sheet_snapshot_date(book, sheet_name, ws) is None:
                    continue
                for row in ws.iter_rows(min_row=5, values_only=True):
                    if not row or row[0] is None or len(row) < 7:
                        continue
                    stock = str(row[0]).strip()
                    change = row[6]
                    if not change or not str(change).strip():
                        continue
                    raw = str(change).strip()
                    if is_liquid_holding(stock):
                        result.skipped.append(
                            f"{book.name}/{sheet_name} {stock}: liquid note skipped ({raw})"
                        )
                        continue
                    if parse_change_note(raw) is None:
                        result.skipped.append(
                            f"{book.name}/{sheet_name} {stock}: unparsed ({raw})"
                        )
        finally:
            wb.close()

    if dry_run or not missing:
        result.appended = []
        return result

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.stem}.BACKUP_{stamp}{path.suffix}")
    shutil.copy2(path, backup)
    result.backup_path = str(backup)

    for trade in missing:
        qty = trade.quantity
        if trade.event_type == EventType.SELL.value:
            qty = -abs(qty)
        _append_transaction_row(
            path,
            {
                "event_date": trade.event_date,
                "event_type": trade.event_type,
                "quantity": qty,
                "price": trade.price,
                "amount": trade.amount,
                "source_note": trade.source_note,
            },
            portfolio_name=trade.portfolio_name,
        )
        result.appended.append(trade)
        fingerprints.add(trade.fingerprint)
        tuples.add(
            (
                trade.portfolio_name.lower(),
                trade.event_date,
                trade.event_type.lower(),
                trade.quantity,
                trade.price,
            )
        )
    synced = _sync_to_raw(MasterKind.TRANSACTIONS, path)
    result.synced_raw = str(synced) if synced is not None else None
    return result
