"""Portfolio snapshot workbook parser and importer."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from pms_platform.ingestion.common import file_checksum, to_decimal, to_int_quantity
from pms_platform.models import ImportBatch, PortfolioSnapshotRecord, Security
from pms_platform.portfolio.name_resolver import build_name_lookup, resolve_snapshot_security_id
from pms_platform.portfolio.types import PortfolioSnapshot

MONTHS = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}

SKIP_SHEETS = frozenset({"liquid", "sheet1", "sheet2", "sheet3", "model"})


@dataclass(frozen=True)
class SnapshotImportResult:
    """Outcome of a portfolio snapshot import."""

    inserted: int
    skipped: int
    unresolved_names: int
    import_batch_id: int


def parse_snapshot_sheet_date(sheet_name: str, workbook_year: int) -> date | None:
    """Parse a snapshot date from a workbook sheet name."""
    cleaned = sheet_name.strip()
    cleaned = cleaned.replace("st", "").replace("nd", "").replace("rd", "").replace("th", "")

    match = re.match(r"^(\d{1,2})([A-Za-z]{3})(\d{2,4})?$", cleaned)
    if match is None:
        return None

    day = int(match.group(1))
    month = MONTHS.get(match.group(2).lower())
    if month is None:
        return None

    year_suffix = match.group(3)
    if year_suffix is None:
        year = workbook_year
    elif len(year_suffix) == 2:
        year = 2000 + int(year_suffix)
    else:
        year = int(year_suffix)
    return date(year, month, day)


def _workbook_year(path: Path) -> int:
    """Extract the year encoded in a portfolio workbook filename."""
    match = re.search(r"Portfolio_(\d{4})", path.name)
    if match is None:
        msg = f"Unable to determine workbook year from {path.name}"
        raise ValueError(msg)
    return int(match.group(1))


def _sheet_snapshot_date(path: Path, sheet_name: str, worksheet) -> date | None:
    """Resolve the snapshot date for a worksheet."""
    from_sheet_name = parse_snapshot_sheet_date(sheet_name, _workbook_year(path))
    if from_sheet_name is not None:
        return from_sheet_name

    rows = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    if rows and len(rows[0]) > 1 and isinstance(rows[0][1], datetime):
        return rows[0][1].date()
    return None


def parse_snapshot_workbook(path: Path) -> list[PortfolioSnapshot]:
    """Parse all snapshot rows from one annual portfolio workbook."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parsed: list[PortfolioSnapshot] = []

    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() in SKIP_SHEETS:
            continue

        worksheet = workbook[sheet_name]
        snapshot_date = _sheet_snapshot_date(path, sheet_name, worksheet)
        if snapshot_date is None:
            continue

        rows = list(worksheet.iter_rows(values_only=True))
        for row_number, row in enumerate(rows, start=1):
            if row_number <= 4:
                continue

            stock_name, quantity, market_price, market_value, portfolio_weight = (
                row + (None,) * 5
            )[:5]
            if stock_name is None or quantity is None:
                continue

            label = str(stock_name).strip()
            if not label or label.lower() in {"total", "cash", "stock"}:
                continue

            parsed.append(
                PortfolioSnapshot(
                    snapshot_date=snapshot_date,
                    portfolio_name=label,
                    quantity=to_int_quantity(quantity),
                    market_price=to_decimal(market_price),
                    market_value=to_decimal(market_value),
                    portfolio_weight=to_decimal(portfolio_weight),
                    source_file=str(path),
                    source_sheet=sheet_name,
                    source_row=row_number,
                )
            )

    workbook.close()
    return parsed


def import_portfolio_snapshots(
    session: Session, directory: Path | None = None
) -> SnapshotImportResult:
    """Import all annual portfolio snapshot workbooks idempotently."""
    snapshot_dir = directory or Path("data/raw/portfolio_snapshots")
    paths = sorted(snapshot_dir.glob("Portfolio_*.xlsx"))
    if not paths:
        msg = f"No portfolio snapshot workbooks found in {snapshot_dir}"
        raise FileNotFoundError(msg)

    combined = hashlib.sha256()
    for path in paths:
        combined.update(file_checksum(path).encode())
    combined_checksum = combined.hexdigest()

    existing_batch = session.scalar(
        select(ImportBatch).where(
            ImportBatch.source_type == "snapshots",
            ImportBatch.source_file == str(snapshot_dir),
            ImportBatch.source_checksum == combined_checksum,
        )
    )
    if existing_batch is not None:
        count = len(
            session.scalars(
                select(PortfolioSnapshotRecord).where(
                    PortfolioSnapshotRecord.import_batch_id == existing_batch.import_batch_id
                )
            ).all()
        )
        return SnapshotImportResult(
            inserted=0,
            skipped=count,
            unresolved_names=0,
            import_batch_id=existing_batch.import_batch_id,
        )

    batch = ImportBatch(
        source_type="snapshots",
        source_file=str(snapshot_dir),
        source_checksum=combined_checksum,
        status="completed",
    )
    session.add(batch)
    session.flush()

    securities = list(session.scalars(select(Security)).all())
    name_lookup = build_name_lookup(session)
    inserted = 0
    skipped = 0
    unresolved_names = 0

    for path in paths:
        for row in parse_snapshot_workbook(path):
            existing = session.scalar(
                select(PortfolioSnapshotRecord).where(
                    PortfolioSnapshotRecord.snapshot_date == row.snapshot_date,
                    PortfolioSnapshotRecord.source_file == row.source_file,
                    PortfolioSnapshotRecord.source_sheet == row.source_sheet,
                    PortfolioSnapshotRecord.source_row == row.source_row,
                )
            )
            if existing is not None:
                skipped += 1
                continue

            security_id, resolved_name = resolve_snapshot_security_id(
                row.portfolio_name,
                securities,
                name_lookup,
            )
            if resolved_name is None:
                unresolved_names += 1
                continue

            session.add(
                PortfolioSnapshotRecord(
                    snapshot_date=row.snapshot_date,
                    security_id=security_id,
                    portfolio_name=resolved_name,
                    quantity=row.quantity,
                    market_price=row.market_price,
                    market_value=row.market_value,
                    portfolio_weight=row.portfolio_weight,
                    source_file=row.source_file,
                    source_sheet=row.source_sheet,
                    source_row=row.source_row,
                    import_batch_id=batch.import_batch_id,
                )
            )
            inserted += 1

    session.flush()
    return SnapshotImportResult(
        inserted=inserted,
        skipped=skipped,
        unresolved_names=unresolved_names,
        import_batch_id=batch.import_batch_id,
    )


def clear_portfolio_snapshots(session: Session) -> None:
    """Delete imported portfolio snapshot rows."""
    session.execute(delete(PortfolioSnapshotRecord))
    session.execute(delete(ImportBatch).where(ImportBatch.source_type == "snapshots"))
