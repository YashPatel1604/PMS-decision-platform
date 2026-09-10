"""Seed pivot portfolio (+ optional history) from DailyEditFiles, else Research."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from pms_platform.market_data.nse_bhav_parse import normalize_bhav_row
from pms_platform.market_data.nse_bhav_store import (
    MAX_BHAV_SESSIONS,
    backfill_vol_exp_snapshots,
    delete_portfolio_symbol,
    prune_bhav_sessions,
    replace_vol_exp_stats,
    sync_bhav_file,
    upsert_portfolio_symbols,
)
from pms_platform.research_paths import research_dir


def default_pivot_workbook() -> Path | None:
    from pms_platform.market_data.daily_edit_bhav import pivot_workbook_daily_edit_path

    hit = pivot_workbook_daily_edit_path()
    if hit is not None:
        return hit
    root = research_dir()
    if root is None:
        return None
    preferred = root / "PivotPointsStrategy_New -Backup 17.07.2024 8PM - Copy.xlsx"
    if preferred.is_file():
        return preferred
    matches = sorted(root.glob("PivotPoints*.xlsx"))
    return matches[0] if matches else None


def _yn(value: object) -> bool:
    return str(value or "").strip().upper() in {"Y", "YES", "TRUE", "1"}


def parse_portfolio_rows(path: Path) -> list[dict[str, Any]]:
    """Read Portfolio sheet rows (no DB writes)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Portfolio" not in wb.sheetnames:
        wb.close()
        raise FileNotFoundError("Workbook has no Portfolio sheet")
    ws = wb["Portfolio"]
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)  # header
    payload: list[dict[str, Any]] = []
    for row in rows_iter:
        if not row or not row[0]:
            continue
        payload.append(
            {
                "symbol": str(row[0]).strip().upper(),
                "dummy": _yn(row[1] if len(row) > 1 else None),
                "portfolio_a": _yn(row[2] if len(row) > 2 else None),
                "uptrend": _yn(row[3] if len(row) > 3 else None),
                "support_note": str(row[4]).strip() if len(row) > 4 and row[4] is not None else None,
                "buy_note": str(row[5]).strip() if len(row) > 5 and row[5] is not None else None,
                "sma_50": str(row[6]).strip() if len(row) > 6 and row[6] is not None else None,
                "sma_100": str(row[7]).strip() if len(row) > 7 and row[7] is not None else None,
                "sma_200": str(row[8]).strip() if len(row) > 8 and row[8] is not None else None,
                "notes": str(row[9]).strip() if len(row) > 9 and row[9] is not None else None,
            }
        )
    wb.close()
    return payload


def seed_portfolio_from_workbook(session: Session, path: Path) -> int:
    return upsert_portfolio_symbols(session, parse_portfolio_rows(path))


def reimport_vol_exp_from_daily(
    session: Session,
    path: Path,
    *,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Replace pivot_vol_exp for the Daily sheet TradDt from Excel 'Vol Exp' column.

    Software-computed Last20 vol often diverges from the workbook; Excel is source
    of truth when the Pivot file is reimported.
    """
    from decimal import Decimal, InvalidOperation

    wb = load_workbook(path, read_only=True, data_only=True)
    if "Daily" not in wb.sheetnames:
        wb.close()
        raise FileNotFoundError("Workbook has no Daily sheet")
    ws = wb["Daily"]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        wb.close()
        raise FileNotFoundError("Daily sheet is empty")
    headers = [str(h).strip() if h is not None else f"c{i}" for i, h in enumerate(header)]
    idx = {h: i for i, h in enumerate(headers)}
    need = ("TradDt", "TckrSymb", "SctySrs", "Vol Exp")
    missing = [c for c in need if c not in idx]
    if missing:
        wb.close()
        raise FileNotFoundError(f"Daily sheet missing columns: {missing}")

    by_sym: dict[str, tuple[Decimal, date]] = {}
    as_of: date | None = None
    for row in rows_iter:
        series = str(row[idx["SctySrs"]] or "").strip().upper()
        if series not in {"EQ", "BE"}:
            continue
        sym = str(row[idx["TckrSymb"]] or "").strip().upper()
        if not sym:
            continue
        raw_dt = row[idx["TradDt"]]
        if hasattr(raw_dt, "date"):
            d = raw_dt.date()
        elif isinstance(raw_dt, date):
            d = raw_dt
        else:
            continue
        raw_vol = row[idx["Vol Exp"]]
        if raw_vol is None or raw_vol == "":
            continue
        try:
            vol = Decimal(str(raw_vol).strip().replace(",", ""))
        except (InvalidOperation, ValueError):
            continue
        as_of = d if as_of is None else max(as_of, d)
        # Prefer EQ when both series present.
        prev = by_sym.get(sym)
        if prev is not None and series != "EQ":
            continue
        by_sym[sym] = (vol, d)
    wb.close()
    if as_of is None or not by_sym:
        return {"vol_exp_symbols": 0, "as_of": None}
    # Keep rows for the latest TradDt only (Daily is usually one session).
    payload = [
        (sym, vol, None)
        for sym, (vol, d) in by_sym.items()
        if d == as_of
    ]
    if dry_run:
        return {"vol_exp_symbols": len(payload), "as_of": as_of.isoformat()}
    n = replace_vol_exp_stats(session, payload, source="excel_daily", as_of=as_of)
    return {"vol_exp_symbols": n, "as_of": as_of.isoformat()}


def reimport_pivot_portfolio(
    session: Session,
    path: Path,
    *,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Upsert Portfolio sheet symbols; remove DB rows missing from Excel.

    Also replaces Vol Exp for the Daily sheet date from Excel (authoritative).
    """
    from sqlalchemy import select

    from pms_platform.models.nse_bhav import PivotPortfolioSymbol

    payload = parse_portfolio_rows(path)
    excel = {str(r["symbol"]).upper() for r in payload if r.get("symbol")}
    existing = set(session.scalars(select(PivotPortfolioSymbol.symbol)).all())
    added = sorted(excel - existing)
    removed = sorted(existing - excel)
    if dry_run:
        vol = reimport_vol_exp_from_daily(session, path, dry_run=True)
        return {
            "added": added,
            "removed": removed,
            "metadata_updated": [],
            "qty_updated": [],
            "portfolio_symbols": len(excel),
            **vol,
        }
    upsert_portfolio_symbols(session, payload)
    for symbol in removed:
        delete_portfolio_symbol(session, symbol)
    vol = reimport_vol_exp_from_daily(session, path, dry_run=False)
    return {
        "added": added,
        "removed": removed,
        "metadata_updated": [],
        "qty_updated": [],
        "portfolio_symbols": len(excel),
        **vol,
    }


def seed_vol_exp_from_all_symbols(session: Session, path: Path, *, as_of: date) -> int:
    """Load Excel AllSymbols col 'AvgQty20Days+x%' as a one-day Vol Exp snapshot."""
    from decimal import Decimal, InvalidOperation

    wb = load_workbook(path, read_only=True, data_only=True)
    if "AllSymbols" not in wb.sheetnames:
        wb.close()
        return 0
    ws = wb["AllSymbols"]
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)
    payload: list[tuple[str, Decimal, int | None]] = []
    for row in rows_iter:
        if not row or not row[0]:
            continue
        symbol = str(row[0]).strip().upper()
        raw = row[4] if len(row) > 4 else None
        if raw is None or raw == "":
            continue
        try:
            vol_exp = Decimal(str(raw).strip().replace(",", ""))
        except (InvalidOperation, ValueError):
            continue
        rank = None
        if len(row) > 3 and row[3] is not None:
            try:
                rank = int(row[3])
            except (TypeError, ValueError):
                rank = None
        payload.append((symbol, vol_exp, rank))
    wb.close()
    return replace_vol_exp_stats(session, payload, source="all_symbols", as_of=as_of)


def _export_sheet_to_csv(path: Path, sheet_name: str, dest: Path) -> int:
    """Write UDiFF-like columns from a workbook sheet to CSV for sync_bhav_file."""
    import csv

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return 0
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        wb.close()
        return 0
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
    # Keep only columns needed by parser (+ extras harmless).
    dest.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with dest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for values in rows_iter:
            if not values or all(v is None or v == "" for v in values):
                continue
            raw = {headers[i]: values[i] for i in range(min(len(headers), len(values)))}
            try:
                if normalize_bhav_row(raw) is None:
                    continue
            except Exception:
                continue
            writer.writerow({k: ("" if raw.get(k) is None else raw.get(k)) for k in headers})
            count += 1
    wb.close()
    return count


def seed_pivot_from_research(
    session: Session,
    *,
    workbook: Path | None = None,
    include_history: bool = True,
) -> dict[str, int]:
    """Seed portfolio symbols; optionally ingest Daily + Last20Days sheets as bhav days."""
    path = workbook or default_pivot_workbook()
    if path is None or not path.is_file():
        raise FileNotFoundError(
            "PivotPoints*.xlsx not found under DailyEditFiles (or Research/)"
        )

    portfolio_n = seed_portfolio_from_workbook(session, path)
    days_committed = 0
    if include_history:
        from tempfile import TemporaryDirectory

        with TemporaryDirectory(prefix="pivot-seed-") as tmp:
            tmp_path = Path(tmp)
            peek = load_workbook(path, read_only=True)
            available = set(peek.sheetnames)
            peek.close()
            # Prefer Last20Days_test (20 working days); fall back to Last20Days.
            history_sheet = next(
                (s for s in ("Last20Days_test", "Last20Days") if s in available),
                None,
            )
            if history_sheet is not None:
                csv_path = tmp_path / f"{history_sheet}.csv"
                n = _export_sheet_to_csv(path, history_sheet, csv_path)
                if n > 0:
                    days_committed += _commit_multiday_csv(session, csv_path)
            if "Daily" in available:
                csv_path = tmp_path / "Daily.csv"
                n = _export_sheet_to_csv(path, "Daily", csv_path)
                if n > 0:
                    run = sync_bhav_file(session, csv_path)
                    if run.status == "committed":
                        days_committed += 1
        prune_bhav_sessions(session, MAX_BHAV_SESSIONS)
    # Rolling Last20 Vol Exp for every kept bhav day (what Daily prints).
    vol_n = backfill_vol_exp_snapshots(session)
    return {
        "portfolio_symbols": portfolio_n,
        "days_committed": days_committed,
        "vol_exp_symbols": vol_n,
    }


def _commit_multiday_csv(session: Session, csv_path: Path) -> int:
    import csv
    from collections import defaultdict

    by_date: dict[str, list[dict[str, str]]] = defaultdict(list)
    with csv_path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames or []
        for row in reader:
            key = str(row.get("TradDt") or "")[:10]
            if key:
                by_date[key].append(row)
    committed = 0
    for day, rows in sorted(by_date.items()):
        day_path = csv_path.with_name(f"bhav_{day}.csv")
        with day_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        run = sync_bhav_file(session, day_path)
        if run.status == "committed":
            committed += 1
    return committed
