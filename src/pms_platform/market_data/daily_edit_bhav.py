"""Keep DailyEditFiles Charts + SCA_LLP in lockstep with committed NSE bhav.

Same store as Pivot / Client Portfolio. Qty stays in Excel; PRICE/VALUE and
cmbhavcopy / BhavCopy_NSE_CM refresh from the day's UDiFF CSV.
"""

from __future__ import annotations

import csv
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from pms_platform.config import settings
from pms_platform.market_data.client_portfolio_parse import _to_decimal
from pms_platform.market_data.nse_bhav_parse import parse_bhav_file
from pms_platform.research_paths import _onedrive_personal_root


def daily_edit_dir() -> Path | None:
    """Writable DailyEditFiles folder (project sibling, not Research)."""
    if os.environ.get("PYTEST_CURRENT_TEST"):
        configured = settings.daily_edit_dir
        if configured is None:
            return None
        path = Path(configured).expanduser().resolve()
        return path if path.is_dir() else None
    if settings.daily_edit_dir is not None:
        path = Path(settings.daily_edit_dir).expanduser().resolve()
        return path if path.is_dir() else None
    cwd = Path.cwd().resolve()
    candidates = (
        cwd / "DailyEditFiles",
        cwd.parent / "DailyEditFiles",
        _onedrive_personal_root() / "PMS-Decision-Platform" / "DailyEditFiles",
        Path("/data/daily_edit"),
    )
    for path in candidates:
        if path.is_dir():
            return path
    return None


def _newest_xlsx(folder: Path, *needles: str) -> Path | None:
    hits: list[Path] = []
    for path in folder.glob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        name = path.name.casefold()
        if all(n.casefold() in name for n in needles):
            hits.append(path)
    if not hits:
        return None
    return max(hits, key=lambda p: p.stat().st_mtime)


def charts_workbook_path(folder: Path | None = None) -> Path | None:
    root = folder or daily_edit_dir()
    return None if root is None else _newest_xlsx(root, "chart")


def sca_llp_workbook_path(folder: Path | None = None) -> Path | None:
    root = folder or daily_edit_dir()
    if root is None:
        return None
    return _newest_xlsx(root, "sca") or _newest_xlsx(root, "llp", "stock")


def client_portfolio_daily_edit_path(folder: Path | None = None) -> Path | None:
    root = folder or daily_edit_dir()
    return None if root is None else _newest_xlsx(root, "pms", "client")


def pivot_workbook_daily_edit_path(folder: Path | None = None) -> Path | None:
    root = folder or daily_edit_dir()
    return None if root is None else _newest_xlsx(root, "pivot")


def _csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            return [], []
        headers = [str(h).strip() for h in reader.fieldnames if h]
        rows = [{str(k).strip(): ("" if v is None else str(v)) for k, v in row.items() if k} for row in reader]
    return headers, rows


def _sheet(wb, *names: str) -> Worksheet | None:
    wanted = {n.casefold() for n in names}
    for title in wb.sheetnames:
        if title.strip().casefold() in wanted:
            return wb[title]
    return None


def _write_udiff_sheet(ws: Worksheet, headers: list[str], rows: list[dict[str, str]]) -> int:
    """Replace data rows; keep header. Charts col A is SYMBOL+Series formula."""
    existing = [cell.value for cell in ws[1]]
    if not existing or existing[0] is None:
        existing = list(headers)
        for col, header in enumerate(existing, start=1):
            ws.cell(1, col, header)
    col_by_name = {str(h).strip(): i for i, h in enumerate(existing, start=1) if h}
    concat_first = str(existing[0] or "").replace(" ", "").upper() in {"SYMBOL+SERIES", "SYMBOLSERIES"}
    tckr_col = col_by_name.get("TckrSymb")
    srs_col = col_by_name.get("SctySrs")

    if ws.max_row and ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    written = 0
    for r_i, row in enumerate(rows, start=2):
        if concat_first and tckr_col and srs_col:
            ws.cell(r_i, 1, f"=CONCATENATE({_col_letter(tckr_col)}{r_i},{_col_letter(srs_col)}{r_i})")
        for header, value in row.items():
            col = col_by_name.get(header)
            if col is None:
                continue
            ws.cell(r_i, col, _excel_cell(value))
        written += 1
    return written


def _col_letter(idx: int) -> str:
    n = idx
    letters = ""
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _excel_cell(raw: str) -> Any:
    if raw == "":
        return None
    if len(raw) >= 10 and raw[4] == "-" and raw[7] == "-":
        try:
            return datetime.fromisoformat(raw[:10])
        except ValueError:
            pass
    try:
        if "." in raw:
            return float(Decimal(raw))
        return int(raw)
    except (InvalidOperation, ValueError):
        return raw


def _close_map(bhav_path: Path) -> dict[str, Decimal]:
    fallback: dict[str, Decimal] = {}
    eq: dict[str, Decimal] = {}
    for row in parse_bhav_file(bhav_path):
        if row.series == "EQ":
            eq[row.symbol] = row.close
        elif row.symbol not in fallback:
            fallback[row.symbol] = row.close
    return {**fallback, **eq}


def _revalue_quantity_sheet(ws: Worksheet, closes: dict[str, Decimal], as_of: date) -> int:
    updated = 0
    header = [cell.value for cell in ws[1]]
    ws.cell(1, 7, as_of.strftime("%d.%m.%Y"))
    for r in range(2, (ws.max_row or 1) + 1):
        symbol = str(ws.cell(r, 1).value or "").strip().upper()
        if not symbol or symbol in {"SYMBOL", "STOCKS"}:
            continue
        close = closes.get(symbol)
        if close is None:
            continue
        qty = _to_decimal(ws.cell(r, 4).value)
        ws.cell(r, 5, float(close))
        if qty is not None:
            ws.cell(r, 6, float(qty * close))
        # G (D−H), H Ramprasath qty, I blocked (=H*E) stay; only G1 date header moves.
        updated += 1
    _ = header
    return updated


def _revalue_stocks_sheet(ws: Worksheet, closes: dict[str, Decimal]) -> int:
    updated = 0
    for r in range(2, (ws.max_row or 1) + 1):
        symbol = str(ws.cell(r, 1).value or "").strip().upper()
        if not symbol or symbol == "SYMBOL":
            continue
        close = closes.get(symbol)
        if close is None:
            continue
        ws.cell(r, 3, float(close))
        updated += 1
    return updated


def apply_bhav_csv_to_daily_edit_files(bhav_path: Path, *, folder: Path | None = None) -> dict[str, int | str | None]:
    """Write committed UDiFF into Charts + SCA_LLP workbooks. No-op if folder missing."""
    root = folder or daily_edit_dir()
    result: dict[str, int | str | None] = {
        "charts_rows": 0,
        "sca_bhav_rows": 0,
        "sca_qty_priced": 0,
        "sca_stocks_priced": 0,
        "charts_file": None,
        "sca_file": None,
        "error": None,
    }
    if root is None:
        result["error"] = "DailyEditFiles not found"
        return result
    path = Path(bhav_path)
    if not path.is_file():
        result["error"] = f"bhav file missing: {path}"
        return result
    headers, rows = _csv_rows(path)
    if not headers or not rows:
        result["error"] = "empty bhav csv"
        return result
    closes = _close_map(path)
    as_of = parse_bhav_file(path)[0].trade_date

    charts = charts_workbook_path(root)
    if charts is not None:
        wb = load_workbook(charts)
        ws = _sheet(wb, "BhavCopy_NSE_CM") or wb[wb.sheetnames[0]]
        result["charts_rows"] = _write_udiff_sheet(ws, headers, rows)
        result["charts_file"] = str(charts)
        wb.save(charts)
        wb.close()

    sca = sca_llp_workbook_path(root)
    if sca is not None:
        wb = load_workbook(sca)
        bhav_ws = _sheet(wb, "cmbhavcopy", "BhavCopy_NSE_CM")
        if bhav_ws is not None:
            result["sca_bhav_rows"] = _write_udiff_sheet(bhav_ws, headers, rows)
        qty_ws = _sheet(wb, "Quantity")
        if qty_ws is not None:
            result["sca_qty_priced"] = _revalue_quantity_sheet(qty_ws, closes, as_of)
        stocks_ws = _sheet(wb, "Stocks")
        if stocks_ws is not None:
            result["sca_stocks_priced"] = _revalue_stocks_sheet(stocks_ws, closes)
        result["sca_file"] = str(sca)
        wb.save(sca)
        wb.close()
        from pms_platform.market_data.client_portfolio_parse import clear_client_portfolio_cache

        clear_client_portfolio_cache()

    return result


def write_sca_bank_balance(amount: Decimal, *, folder: Path | None = None) -> float:
    """Write Quantity!F on the Balance with Bank row. Creates the row if missing."""
    sca = sca_llp_workbook_path(folder)
    if sca is None:
        raise FileNotFoundError("SCA_LLP Stock Holding.xlsx not found")
    value = float(amount)
    wb = load_workbook(sca)
    ws = _sheet(wb, "Quantity")
    if ws is None:
        wb.close()
        raise FileNotFoundError("Quantity sheet missing")
    row_i = None
    for r in range(1, (ws.max_row or 1) + 1):
        if str(ws.cell(r, 1).value or "").strip().lower() == "balance with bank":
            row_i = r
            break
    if row_i is None:
        row_i = (ws.max_row or 1) + 1
        ws.cell(row_i, 1, "Balance with Bank")
    ws.cell(row_i, 6, value)
    wb.save(sca)
    wb.close()
    from pms_platform.market_data.client_portfolio_parse import clear_client_portfolio_cache

    clear_client_portfolio_cache()
    return value


def maybe_apply_committed_bhav(staged_path: str | None) -> dict[str, int | str | None] | None:
    """Called after a successful bhav commit. Swallows IO errors so ingest never rolls back."""
    if not staged_path:
        return None
    try:
        return apply_bhav_csv_to_daily_edit_files(Path(staged_path))
    except Exception as exc:  # noqa: BLE001
        return {"error": str(exc), "charts_rows": 0, "sca_bhav_rows": 0}
