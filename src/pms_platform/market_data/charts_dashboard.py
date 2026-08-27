"""Charts Range sheet: High/Low Fibonacci targets × as-of bhav close."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from pms_platform.domain.charts_rows import merge_db_into_parsed, sync_rows_from_parse
from pms_platform.feature_flags import approval_workflow_enabled
from pms_platform.market_data.client_portfolio_parse import _to_decimal
from pms_platform.market_data.daily_edit_bhav import charts_workbook_path
from pms_platform.market_data.nse_bhav_store import available_trade_dates, load_day_bars

# Range columns (1-based)
_COL_HIGH = 2
_COL_LOW = 3
_COL_CLOSE = 13
_COL_WEEKLY_CLOSE = 19
_COL_SR = 20
_COL_WEEKLY_DATE = 21

_TRG = (
    ("trg_13", Decimal("0.13")),
    ("trg_21", Decimal("0.21")),
    ("trg_34", Decimal("0.34")),
    ("trg_55", Decimal("0.55")),
    ("trg_89", Decimal("0.89")),
)


def _f(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _series_from_lookup(raw: object) -> str:
    text = str(raw or "").upper()
    return "BE" if '"BE"' in text else "EQ"


def _norm_label(value: str) -> str:
    return " ".join(value.casefold().split())


def _cell_text(raw: object) -> str | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    text = str(raw).strip()
    return text or None


def _numeric_cell(raw: object) -> Decimal | None:
    """Read a typed number; ignore Excel formulas (use bhav instead)."""
    if raw is None or raw == "":
        return None
    if isinstance(raw, str) and raw.strip().startswith("="):
        return None
    return _to_decimal(raw)


def parse_charts_range(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = next((wb[n] for n in wb.sheetnames if n.strip().casefold() == "range"), None)
        if ws is None:
            return []
        rows: list[dict[str, Any]] = []
        section = "holdings"
        for excel_row, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            name = str(row[0]).strip() if row and row[0] not in (None, "") else ""
            if not name:
                continue
            if _norm_label(name) == "nifty fno stocks":
                section = "fno"
                continue
            high = _to_decimal(row[1] if len(row) > 1 else None)
            low = _to_decimal(row[2] if len(row) > 2 else None)
            rows.append(
                {
                    "name": name,
                    "symbol": name.upper().replace(" ", ""),
                    "section": section,
                    "excel_row": excel_row,
                    "series": _series_from_lookup(row[12] if len(row) > 12 else None),
                    "high": high,
                    "low": low,
                    "excel_close": _numeric_cell(row[12] if len(row) > 12 else None),
                    "weekly_close": _to_decimal(row[18] if len(row) > 18 else None),
                    "support_resistance": _cell_text(row[19] if len(row) > 19 else None),
                    "weekly_close_date": _cell_text(row[20] if len(row) > 20 else None),
                }
            )
        return rows
    finally:
        wb.close()


def _levels(high: Decimal | None, low: Decimal | None) -> dict[str, Decimal | None]:
    if high is None or low is None:
        return {key: None for key, _ in _TRG} | {"difference": None, "trg_144": None}
    diff = high - low
    out: dict[str, Decimal | None] = {"difference": diff}
    for key, frac in _TRG:
        out[key] = diff * frac + low
    # Excel J = E*144% + D
    trg_13 = out["trg_13"]
    out["trg_144"] = (trg_13 * Decimal("1.44") + diff) if trg_13 is not None else None
    return out


def _pick_bar(eq: dict[str, Any], be: dict[str, Any], symbol: str, series: str) -> Any | None:
    if series == "BE":
        return be.get(symbol) or eq.get(symbol)
    return eq.get(symbol) or be.get(symbol)


def build_charts_dashboard(
    session: Session,
    *,
    as_of: date | None = None,
) -> dict[str, Any]:
    dates = available_trade_dates(session)
    if as_of is None:
        as_of = dates[0] if dates else None
    path = charts_workbook_path()
    empty = {
        "as_of": as_of.isoformat() if as_of else None,
        "available_dates": [d.isoformat() for d in dates],
        "source_file": str(path) if path else None,
        "excel_mtime": None,
        "rows": [],
        "missing_symbols": [],
        "error": None,
    }
    if path is None or not path.is_file():
        empty["error"] = (
            "Charts.xlsx not found under DailyEditFiles "
            "(set DAILY_EDIT_DIR / pin Always keep on this device)."
        )
        return empty
    empty["excel_mtime"] = datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    names = parse_charts_range(path)
    workflow = approval_workflow_enabled()
    if workflow:
        sync_rows_from_parse(session, names)
        session.flush()
        names = merge_db_into_parsed(session, names)
    if as_of is None:
        empty["error"] = "No bhav days committed yet — upload on Pivot Point."
        return empty

    eq = {b.symbol: b for b in load_day_bars(session, as_of, series="EQ")}
    be = {b.symbol: b for b in load_day_bars(session, as_of, series="BE")}
    missing: list[str] = []
    out_rows: list[dict[str, Any]] = []
    for raw in names:
        levels = _levels(raw["high"], raw["low"])
        bar = _pick_bar(eq, be, raw["symbol"], raw["series"])
        bhav_close = Decimal(bar.close) if bar is not None else None
        # Typed Excel close overrides bhav (after user edit); formulas still use bhav.
        close = raw.get("excel_close") if raw.get("excel_close") is not None else bhav_close
        prev = Decimal(bar.prev_close) if bar is not None and bar.prev_close is not None else None
        if bar is None and raw.get("excel_close") is None:
            missing.append(raw["symbol"])
        pct_from_lows = None
        if close is not None and raw["low"] not in (None, Decimal(0)):
            pct_from_lows = (close * Decimal(100) / raw["low"]) - Decimal(100)
        trg_89 = levels["trg_89"]
        high = raw["high"]
        low = raw["low"]
        row = {
            "name": raw["name"],
            "symbol": raw["symbol"],
            "section": raw["section"],
            "excel_row": raw["excel_row"],
            "series": raw["series"] if bar is None else bar.series,
            "high": _f(high),
            "low": _f(low),
            "difference": _f(levels["difference"]),
            "trg_13": _f(levels["trg_13"]),
            "trg_21": _f(levels["trg_21"]),
            "trg_34": _f(levels["trg_34"]),
            "trg_55": _f(levels["trg_55"]),
            "trg_89": _f(levels["trg_89"]),
            "trg_144": _f(levels["trg_144"]),
            "close": _f(close),
            "bhav_close": _f(bhav_close),
            "prev_close": _f(prev),
            "pct_from_lows": _f(pct_from_lows),
            "corr_10": _f(high - close * Decimal("0.10"))
            if close is not None and high is not None
            else None,
            "corr_20": _f(high - close * Decimal("0.20"))
            if close is not None and high is not None
            else None,
            "below_trg_89": bool(close is not None and trg_89 is not None and close < trg_89),
            "below_low": bool(close is not None and low is not None and close < low),
            "above_high": bool(close is not None and high is not None and close > high),
            "weekly_close": _f(raw["weekly_close"]),
            "support_resistance": raw["support_resistance"],
            "weekly_close_date": raw["weekly_close_date"],
            "missing_bhav": bar is None and raw.get("excel_close") is None,
        }
        out_rows.append(row)
    empty["rows"] = out_rows
    empty["missing_symbols"] = missing
    return empty


def write_charts_range_hlc(
    excel_row: int,
    *,
    high: float | None = None,
    low: float | None = None,
    close: float | None = None,
    folder: Path | None = None,
) -> dict[str, Any]:
    """Write Range High / Low / Close (cols B/C/M)."""
    if high is None and low is None and close is None:
        raise ValueError("Provide high, low, and/or close")
    path = charts_workbook_path(folder)
    if path is None or not path.is_file():
        raise FileNotFoundError("Charts.xlsx not found")
    wb = load_workbook(path)
    try:
        ws = next((wb[n] for n in wb.sheetnames if n.strip().casefold() == "range"), None)
        if ws is None:
            raise FileNotFoundError("Range sheet missing")
        name = str(ws.cell(excel_row, 1).value or "").strip()
        if not name or _norm_label(name) == "nifty fno stocks":
            raise ValueError("That Excel row is not a stock.")

        def _num(label: str, raw: float | None) -> float | None:
            if raw is None:
                return None
            try:
                return float(Decimal(str(raw)))
            except (InvalidOperation, ValueError) as exc:
                raise ValueError(f"{label} must be a number") from exc

        high_v = _num("high", high)
        low_v = _num("low", low)
        close_v = _num("close", close)
        if high_v is not None:
            ws.cell(excel_row, _COL_HIGH, high_v)
        if low_v is not None:
            ws.cell(excel_row, _COL_LOW, low_v)
        if close_v is not None:
            ws.cell(excel_row, _COL_CLOSE, close_v)
        wb.save(path)
        return {
            "excel_row": excel_row,
            "high": high_v if high_v is not None else _f(_to_decimal(ws.cell(excel_row, _COL_HIGH).value)),
            "low": low_v if low_v is not None else _f(_to_decimal(ws.cell(excel_row, _COL_LOW).value)),
            "close": close_v
            if close_v is not None
            else _f(_numeric_cell(ws.cell(excel_row, _COL_CLOSE).value)),
        }
    finally:
        wb.close()


def write_charts_range_weekly(
    excel_row: int,
    *,
    weekly_close: float | None,
    support_resistance: str | None,
    weekly_close_date: str | None,
    folder: Path | None = None,
) -> dict[str, Any]:
    """Write Range S/T/U on a holdings row. FNO rows are not writable."""
    path = charts_workbook_path(folder)
    if path is None or not path.is_file():
        raise FileNotFoundError("Charts.xlsx not found")
    wb = load_workbook(path)
    try:
        ws = next((wb[n] for n in wb.sheetnames if n.strip().casefold() == "range"), None)
        if ws is None:
            raise FileNotFoundError("Range sheet missing")
        section = "holdings"
        for r in range(3, excel_row + 1):
            label = str(ws.cell(r, 1).value or "").strip()
            if _norm_label(label) == "nifty fno stocks":
                section = "fno"
                break
        if section != "holdings":
            raise ValueError("Weekly S/R is only editable on holdings.")
        name = str(ws.cell(excel_row, 1).value or "").strip()
        if not name or _norm_label(name) == "nifty fno stocks":
            raise ValueError("That Excel row is not a holding.")
        close_val: float | None = None
        if weekly_close is not None:
            try:
                close_val = float(Decimal(str(weekly_close)))
            except (InvalidOperation, ValueError) as exc:
                raise ValueError("weekly_close must be a number") from exc
        ws.cell(excel_row, _COL_WEEKLY_CLOSE, close_val)
        ws.cell(excel_row, _COL_SR, (support_resistance or "").strip() or None)
        ws.cell(excel_row, _COL_WEEKLY_DATE, (weekly_close_date or "").strip() or None)
        wb.save(path)
        return {
            "excel_row": excel_row,
            "weekly_close": close_val,
            "support_resistance": (support_resistance or "").strip() or None,
            "weekly_close_date": (weekly_close_date or "").strip() or None,
        }
    finally:
        wb.close()
