"""Build DailyEdit workbook bytes from live software state (DB + bhav).

Backup must snapshot what the app shows, not a stale DailyEditFiles copy.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy.orm import Session


def _wb_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def _load(data: bytes):
    return load_workbook(io.BytesIO(data))


def _patch_sca(session: Session, data: bytes) -> bytes:
    from pms_platform.market_data.client_portfolio_dashboard import build_client_portfolio_dashboard

    dash = build_client_portfolio_dashboard(session, book="sca")
    by_sym = {str(r["symbol"]).upper(): r for r in dash.get("holdings") or []}
    as_of = dash.get("as_of")
    bank = dash.get("bank_balance")
    total = dash.get("total_value")
    portfolio = dash.get("portfolio_total")

    wb = _load(data)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().casefold() == "quantity"), None)
    if ws is None:
        wb.close()
        return data
    if as_of:
        ws.cell(1, 7, date.fromisoformat(as_of).strftime("%d.%m.%Y"))
    equity_end = None
    for r in range(2, (ws.max_row or 1) + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        lower = label.lower().replace(" ", "_")
        if lower == "balance_with_bank":
            if bank is not None:
                ws.cell(r, 6, float(bank))
            continue
        if not label:
            # Total equity row (=SUM…) and Total Portfolio row — stamp live numbers.
            val = ws.cell(r, 6).value
            if isinstance(val, str) and val.upper().startswith("=SUM(") and total is not None:
                ws.cell(r, 6, float(total))
                equity_end = r
            elif (
                equity_end is not None
                and r > equity_end
                and isinstance(val, str)
                and val.startswith("=")
                and portfolio is not None
            ):
                ws.cell(r, 6, float(portfolio))
            continue
        row = by_sym.get(label.upper())
        if row is None:
            continue
        if row.get("qty") is not None:
            ws.cell(r, 4, float(row["qty"]))
        if row.get("price") is not None:
            ws.cell(r, 5, float(row["price"]))
        if row.get("value") is not None:
            ws.cell(r, 6, float(row["value"]))
        # Percent of book when we have total.
        if total and row.get("value") is not None and total != 0:
            ws.cell(r, 3, float(row["value"]) / float(total))
    return _wb_bytes(wb)


def _patch_charts(session: Session, data: bytes) -> bytes:
    from sqlalchemy import select

    from pms_platform.models.charts_range import ChartsRangeRow

    rows = list(session.scalars(select(ChartsRangeRow)).all())
    if not rows:
        return data
    by_excel = {r.excel_row: r for r in rows}
    wb = _load(data)
    ws = next((wb[n] for n in wb.sheetnames if n.strip().casefold() == "range"), None)
    if ws is None:
        wb.close()
        return data
    for excel_row, row in by_excel.items():
        if row.high is not None:
            ws.cell(excel_row, 2, float(row.high))
        if row.low is not None:
            ws.cell(excel_row, 3, float(row.low))
        if row.close_override is not None:
            ws.cell(excel_row, 13, float(row.close_override))
        if row.weekly_close is not None:
            ws.cell(excel_row, 19, float(row.weekly_close))
        if row.support_resistance is not None:
            ws.cell(excel_row, 20, row.support_resistance)
        if row.weekly_close_date is not None:
            ws.cell(excel_row, 21, row.weekly_close_date)
    return _wb_bytes(wb)


def _patch_client(session: Session, data: bytes) -> bytes:
    from pms_platform.market_data.client_portfolio_dashboard import build_client_portfolio_dashboard

    dash = build_client_portfolio_dashboard(session, book="client")
    by_sym = {str(r["symbol"]).upper(): r for r in dash.get("holdings") or []}
    total = dash.get("total_value")
    wb = _load(data)
    if "Model" not in wb.sheetnames:
        wb.close()
        return data
    ws = wb["Model"]
    for r in range(1, (ws.max_row or 1) + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        if not label:
            continue
        if label.lower().replace(" ", "_") in {"total_value", "total"}:
            if total is not None:
                ws.cell(r, 4, float(total))
                ws.cell(r, 5, 100)
            continue
        row = by_sym.get(label.upper())
        if row is None:
            continue
        if row.get("qty") is not None:
            ws.cell(r, 2, float(row["qty"]))
        if row.get("price") is not None:
            ws.cell(r, 3, float(row["price"]))
        if row.get("value") is not None:
            ws.cell(r, 4, float(row["value"]))
        if row.get("percent") is not None:
            ws.cell(r, 5, float(row["percent"]))
        if row.get("mcap") is not None:
            ws.cell(r, 7, float(row["mcap"]))
        if row.get("firm_pct") is not None:
            ws.cell(r, 9, float(row["firm_pct"]))
    return _wb_bytes(wb)


def _as_date(raw: object) -> date | None:
    if raw is None:
        return None
    if hasattr(raw, "date") and callable(raw.date):
        try:
            return raw.date()  # type: ignore[no-any-return]
        except Exception:  # noqa: BLE001
            pass
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    if len(text) >= 10 and text[4] == "-":
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            return None
    return None


def _trim_pivot_history_sheets(wb, keep: int = 21) -> None:
    """Keep only the newest ``keep`` TradDt days on Last20Days* sheets."""
    for name in ("Last20Days_test", "Last20Days"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        header, body = rows[0], rows[1:]
        dates = sorted(
            {d for d in (_as_date(r[0]) for r in body) if d is not None},
            reverse=True,
        )
        keep_set = set(dates[:keep])
        if not keep_set or len(dates) <= keep:
            continue
        kept = [header] + [r for r in body if _as_date(r[0]) in keep_set]
        ws.delete_rows(1, ws.max_row)
        for r_i, row in enumerate(kept, start=1):
            for c_i, val in enumerate(row, start=1):
                if val is not None:
                    ws.cell(r_i, c_i, val)


def _patch_pivot(session: Session, data: bytes) -> bytes:
    from sqlalchemy import select

    from pms_platform.market_data.nse_bhav_store import (
        MAX_BHAV_SESSIONS,
        available_trade_dates,
        load_day_bars,
        load_vol_exp_map,
    )
    from pms_platform.market_data.pivot_derived import floor_pivot_levels
    from pms_platform.models.nse_bhav import PivotPortfolioSymbol

    dates = available_trade_dates(session)
    if not dates:
        return data
    as_of = dates[0]
    vol = load_vol_exp_map(session, as_of)
    eq = {b.symbol: b for b in load_day_bars(session, as_of, series="EQ")}
    be = {b.symbol: b for b in load_day_bars(session, as_of, series="BE")}
    portfolio = {
        r.symbol: r
        for r in session.scalars(select(PivotPortfolioSymbol)).all()
    }

    wb = _load(data)
    _trim_pivot_history_sheets(wb, keep=MAX_BHAV_SESSIONS)
    if "Portfolio" in wb.sheetnames and portfolio:
        ws = wb["Portfolio"]
        for r in range(2, (ws.max_row or 1) + 1):
            sym = str(ws.cell(r, 1).value or "").strip().upper()
            row = portfolio.get(sym)
            if row is None:
                continue
            ws.cell(r, 2, "Y" if row.dummy else "N")
            ws.cell(r, 3, "Y" if row.portfolio_a else "N")
            ws.cell(r, 4, "Y" if row.uptrend else "N")

    if "Daily" in wb.sheetnames:
        ws = wb["Daily"]
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        idx = {h: i + 1 for i, h in enumerate(headers)}
        c_sym = idx.get("TckrSymb")
        c_series = idx.get("SctySrs")
        c_o, c_h, c_l, c_c = idx.get("OpnPric"), idx.get("HghPric"), idx.get("LwPric"), idx.get("ClsPric")
        c_vol = idx.get("Vol Exp")
        c_15 = idx.get("15minVol")
        c_piv = idx.get("Pivot")
        c_s1, c_r1 = idx.get("S1"), idx.get("R1")
        if c_sym:
            for r in range(2, (ws.max_row or 1) + 1):
                sym = str(ws.cell(r, c_sym).value or "").strip().upper()
                if not sym:
                    continue
                series = str(ws.cell(r, c_series).value or "").strip().upper() if c_series else "EQ"
                if series not in {"EQ", "BE"}:
                    continue
                bar = eq.get(sym) if series == "EQ" else (be.get(sym) or eq.get(sym))
                if bar is None:
                    bar = eq.get(sym) or be.get(sym)
                if bar is not None:
                    if c_o:
                        ws.cell(r, c_o, float(bar.open))
                    if c_h:
                        ws.cell(r, c_h, float(bar.high))
                    if c_l:
                        ws.cell(r, c_l, float(bar.low))
                    if c_c:
                        ws.cell(r, c_c, float(bar.close))
                    levels = floor_pivot_levels(bar.high, bar.low, bar.close)
                    if c_piv:
                        ws.cell(r, c_piv, float(levels["pp"]))
                    if c_s1:
                        ws.cell(r, c_s1, float(levels["s1"]))
                    if c_r1:
                        ws.cell(r, c_r1, float(levels["r1"]))
                ve = vol.get(sym)
                if ve is not None:
                    if c_vol:
                        ws.cell(r, c_vol, float(ve))
                    if c_15:
                        ws.cell(r, c_15, float(Decimal(ve) / Decimal(25)))
    return _wb_bytes(wb)


_PATCHERS = {
    "sca_llp": _patch_sca,
    "charts": _patch_charts,
    "client_portfolio": _patch_client,
    "pivot_points": _patch_pivot,
}


def patch_live_workbook(session: Session, category: str, data: bytes) -> bytes:
    """Overlay live software numbers onto a workbook template."""
    patcher = _PATCHERS.get(category)
    if patcher is None:
        return data
    try:
        return patcher(session, data)
    except Exception:  # noqa: BLE001 — backup should still ship template if patch fails
        return data
