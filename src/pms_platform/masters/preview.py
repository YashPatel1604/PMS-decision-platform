"""Paginated Excel preview for Final Master workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from pms_platform.masters.paths import MasterKind, default_sheet, resolve_master_path


def _cell_to_json(value: Any) -> str | int | float | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()  # type: ignore[no-any-return]
        except Exception:  # noqa: BLE001
            return str(value)
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    return text if text else None


@dataclass(frozen=True)
class WorkbookPreview:
    kind: MasterKind
    path: str
    sheet: str
    sheets: list[str]
    columns: list[str]
    rows: list[dict[str, str | int | float | None]]
    offset: int
    limit: int
    total_rows: int
    matched_rows: int


def preview_workbook(
    kind: MasterKind,
    *,
    sheet: str | None = None,
    offset: int = 0,
    limit: int = 50,
    q: str | None = None,
    master_dir: Path | None = None,
) -> WorkbookPreview:
    path = resolve_master_path(kind, master_dir)
    if not path.is_file():
        msg = f"Workbook not found: {path}"
        raise FileNotFoundError(msg)

    sheet_name = sheet or default_sheet(kind)
    offset = max(0, offset)
    limit = max(1, min(limit, 500))

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = list(workbook.sheetnames)
        if sheet_name not in workbook.sheetnames:
            msg = f"Sheet {sheet_name!r} not in {path.name}; available: {sheets}"
            raise ValueError(msg)
        worksheet = workbook[sheet_name]
        raw_rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not raw_rows:
        return WorkbookPreview(
            kind=kind,
            path=str(path),
            sheet=sheet_name,
            sheets=sheets,
            columns=[],
            rows=[],
            offset=0,
            limit=limit,
            total_rows=0,
            matched_rows=0,
        )

    header = [_cell_to_json(c) for c in raw_rows[0]]
    columns: list[str] = []
    for idx, cell in enumerate(header):
        if cell is None or cell == "":
            columns.append(f"col_{idx + 1}")
        else:
            columns.append(str(cell))

    query = (q or "").strip().lower()
    matched: list[dict[str, str | int | float | None]] = []
    for row_number, row in enumerate(raw_rows[1:], start=2):
        values = list(row) + [None] * max(0, len(columns) - len(row or ()))
        record: dict[str, str | int | float | None] = {"_row": row_number}
        for col_idx, col_name in enumerate(columns):
            record[col_name] = _cell_to_json(values[col_idx] if col_idx < len(values) else None)
        if query:
            haystack = " ".join(
                str(v).lower() for v in record.values() if v is not None and v != ""
            )
            if query not in haystack:
                continue
        matched.append(record)

    page = matched[offset : offset + limit]
    return WorkbookPreview(
        kind=kind,
        path=str(path),
        sheet=sheet_name,
        sheets=sheets,
        columns=columns,
        rows=page,
        offset=offset,
        limit=limit,
        total_rows=max(0, len(raw_rows) - 1),
        matched_rows=len(matched),
    )
