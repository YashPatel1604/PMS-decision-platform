"""Import Screener.in screen/watchlist CSV/XLSX exports into snapshot tables."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from pms_platform.fundamentals.catalog import COMPUTATION_VERSION
from pms_platform.fundamentals.providers.annual_xbrl import upsert_annual_snapshot
from pms_platform.fundamentals.providers.promoter import _upsert_promoter_snapshot
from pms_platform.fundamentals.providers.valuation import upsert_valuation_snapshot
from pms_platform.market_data.identifiers import IdentifierResolver
from pms_platform.models.fundamental_snapshot import FundamentalSnapshot

# Normalize Screener header → our field. First match wins.
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "bse_code": ("bse code", "bse", "bsecode", "security code"),
    "nse_symbol": ("nse code", "nse", "nsecode", "symbol"),
    "name": ("name", "company", "company name", "stock name"),
    "market_cap_cr": (
        "mar cap rs.cr.",
        "mar cap rs.cr",
        "market capitalization",
        "market cap",
        "mar cap",
        "mcap",
    ),
    "pe_ratio": ("p/e", "pe", "price to earning", "price to earnings", "price earning"),
    "industry_pe": ("industry pe", "industry p/e", "ind pe"),
    "price_to_book": ("p/b", "pb", "price to book", "price/book"),
    "price_to_sales": ("p/s", "ps", "price to sales", "price/sales"),
    "eps": ("eps rs.", "eps rs", "eps", "earning per share", "earnings per share"),
    "book_value_per_share": ("book value rs.", "book value rs", "book value", "bv"),
    "dividend_yield": ("div yld %", "div yld", "dividend yield", "div yield"),
    "peg_ratio": ("peg", "peg ratio"),
    "week_52_high": ("52w high", "52 week high", "high 52w"),
    "week_52_low": ("52w low", "52 week low", "low 52w"),
    "return_1m_pct": ("return over 1month", "1m return", "return 1m"),
    "return_3m_pct": ("return over 3months", "3m return", "return 3m"),
    "return_6m_pct": ("return over 6months", "6m return", "return 6m"),
    "return_1y_pct": ("return over 1year", "1y return", "return 1y"),
    "return_3y_pct": ("return over 3years", "3y return", "return 3y"),
    "sales": ("sales qtr rs.cr.", "sales qtr rs.cr", "sales qtr", "sales"),
    "pat": ("np qtr rs.cr.", "np qtr rs.cr", "np qtr", "pat", "profit qtr"),
    "opm": ("opm %", "opm", "operating profit margin"),
    "npm": ("npm %", "npm", "net profit margin"),
    "sales_yoy_pct": ("qtr sales var %", "sales var %", "sales yoy", "sales growth"),
    "pat_yoy_pct": ("qtr profit var %", "profit var %", "pat yoy", "profit growth"),
    "sales_3y_cagr": ("sales growth 3years", "sales 3years", "sales cagr 3y"),
    "sales_5y_cagr": ("sales growth 5years", "sales 5years", "sales cagr 5y"),
    "pat_3y_cagr": ("profit growth 3years", "profit 3years", "pat cagr 3y"),
    "pat_5y_cagr": ("profit growth 5years", "profit 5years", "pat cagr 5y"),
    "roce": ("roce %", "roce", "return on capital employed"),
    "roe": ("roe %", "roe", "return on equity"),
    "roa": ("roa %", "roa", "return on assets"),
    "debt_to_equity": ("debt / eq", "debt/eq", "debt to equity", "d/e", "de"),
    "interest_coverage": ("interest cover", "int. coverage", "interest coverage"),
    "current_ratio": ("current ratio", "curr ratio"),
    "promoter_holding_pct": ("promoter holding", "promoter hold %", "promoter %"),
    "pledged_pct": ("pledged %", "pledge %", "promoter pledging"),
}


@dataclass(frozen=True)
class ScreenerExportImportResult:
    path: Path
    rows_read: int
    valuation_upserts: int
    fund_upserts: int
    annual_upserts: int
    promoter_upserts: int
    skipped: int
    bse_codes: tuple[str, ...]


def _norm_header(raw: str) -> str:
    text = str(raw or "").strip().lower()
    text = text.replace("₹", "rs").replace("%", " %")
    text = re.sub(r"\s+", " ", text)
    return text.strip(" .")


def _build_header_map(fieldnames: list[str]) -> dict[str, str]:
    """Map our field → actual CSV header present in file."""
    by_norm = {_norm_header(h): h for h in fieldnames if h}
    out: dict[str, str] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            hit = by_norm.get(_norm_header(alias))
            if hit:
                out[field] = hit
                break
    return out


def _to_dec(raw: object) -> Decimal | None:
    if raw is None:
        return None
    text = str(raw).strip().replace(",", "").replace("%", "")
    if not text or text.lower() in {"nan", "none", "null", "-", "—", "--", "na", "n/a"}:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _quarter_end(today: date | None = None) -> date:
    """Most recent calendar quarter-end (Mar/Jun/Sep/Dec)."""
    d = today or date.today()
    ends = [
        date(y, m, day)
        for y in (d.year, d.year - 1)
        for m, day in ((3, 31), (6, 30), (9, 30), (12, 31))
    ]
    return max(e for e in ends if e <= d)


def _fiscal_from_period_end(period_end: date) -> tuple[int, str]:
    # Indian FY: Apr–Mar; Q1=Apr-Jun … Q4=Jan-Mar
    if period_end.month in (4, 5, 6):
        return period_end.year + 1, "Q1"
    if period_end.month in (7, 8, 9):
        return period_end.year + 1, "Q2"
    if period_end.month in (10, 11, 12):
        return period_end.year + 1, "Q3"
    return period_end.year, "Q4"


def find_latest_screener_export(directory: Path) -> Path | None:
    """Newest .csv / .xlsx in directory (non-recursive)."""
    if not directory.is_dir():
        return None
    files = [
        p
        for p in directory.iterdir()
        if p.is_file() and p.suffix.lower() in {".csv", ".xlsx", ".xls"} and not p.name.startswith("~")
    ]
    if not files:
        return None
    return max(files, key=lambda p: p.stat().st_mtime)


def _read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            if not reader.fieldnames:
                return [], []
            fields = list(reader.fieldnames)
            return fields, list(reader)
    if suffix in {".xlsx", ".xls"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return [], []
        fields = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(header_row)]
        out: list[dict[str, Any]] = []
        for row in rows_iter:
            out.append({fields[i]: row[i] if i < len(row) else None for i in range(len(fields))})
        return fields, out
    return [], []


def _cell(row: dict[str, Any], header_map: dict[str, str], field: str) -> object:
    header = header_map.get(field)
    if not header:
        return None
    return row.get(header)


def _upsert_fund_snapshot(
    session: Session,
    *,
    identifier_type: str,
    identifier: str,
    security_id: str | None,
    period_end: date,
    sales: Decimal | None,
    pat: Decimal | None,
    opm: Decimal | None,
    npm: Decimal | None,
    sales_yoy_pct: Decimal | None,
    pat_yoy_pct: Decimal | None,
    sales_3y_cagr: Decimal | None,
    sales_5y_cagr: Decimal | None,
    pat_3y_cagr: Decimal | None,
    pat_5y_cagr: Decimal | None,
) -> bool:
    if all(
        v is None
        for v in (
            sales,
            pat,
            opm,
            npm,
            sales_yoy_pct,
            pat_yoy_pct,
            sales_3y_cagr,
            sales_5y_cagr,
            pat_3y_cagr,
            pat_5y_cagr,
        )
    ):
        return False
    fy, fq = _fiscal_from_period_end(period_end)
    now = datetime.now(timezone.utc)
    existing = session.scalar(
        select(FundamentalSnapshot).where(
            FundamentalSnapshot.identifier_type == identifier_type,
            FundamentalSnapshot.identifier == identifier,
            FundamentalSnapshot.period_end_date == period_end,
            FundamentalSnapshot.computation_version == COMPUTATION_VERSION,
        )
    )
    if existing is not None:
        existing.security_id = security_id or existing.security_id
        existing.fiscal_year = fy
        existing.fiscal_quarter = fq
        existing.provider = "screener"
        existing.retrieved_at = now
        existing.computed_at = now
        for key, val in (
            ("sales", sales),
            ("pat", pat),
            ("opm", opm),
            ("npm", npm),
            ("sales_yoy_pct", sales_yoy_pct),
            ("pat_yoy_pct", pat_yoy_pct),
            ("sales_3y_cagr", sales_3y_cagr),
            ("sales_5y_cagr", sales_5y_cagr),
            ("pat_3y_cagr", pat_3y_cagr),
            ("pat_5y_cagr", pat_5y_cagr),
        ):
            if val is not None:
                setattr(existing, key, val)
        return False
    session.add(
        FundamentalSnapshot(
            identifier_type=identifier_type,
            identifier=identifier,
            period_end_date=period_end,
            computation_version=COMPUTATION_VERSION,
            security_id=security_id,
            fiscal_year=fy,
            fiscal_quarter=fq,
            sales=sales,
            pat=pat,
            opm=opm,
            npm=npm,
            sales_yoy_pct=sales_yoy_pct,
            pat_yoy_pct=pat_yoy_pct,
            sales_3y_cagr=sales_3y_cagr,
            sales_5y_cagr=sales_5y_cagr,
            pat_3y_cagr=pat_3y_cagr,
            pat_5y_cagr=pat_5y_cagr,
            provider="screener",
            retrieved_at=now,
            computed_at=now,
        )
    )
    return True


def import_screener_export(
    session: Session,
    path: Path,
    *,
    as_of: date | None = None,
) -> ScreenerExportImportResult:
    """Parse one Screener export and upsert valuation / fund / annual / promoter snapshots."""
    fieldnames, rows = _read_rows(path)
    if not fieldnames or not rows:
        return ScreenerExportImportResult(
            path=path,
            rows_read=0,
            valuation_upserts=0,
            fund_upserts=0,
            annual_upserts=0,
            promoter_upserts=0,
            skipped=0,
            bse_codes=(),
        )

    header_map = _build_header_map(fieldnames)
    if "bse_code" not in header_map and "nse_symbol" not in header_map:
        raise ValueError(
            f"Screener export {path.name} needs a BSE Code or NSE Code column "
            f"(got: {fieldnames[:12]}…)"
        )

    today = as_of or date.today()
    period_end = _quarter_end(today)
    resolver = IdentifierResolver(session)
    val_n = fund_n = ann_n = prom_n = skipped = 0
    codes: list[str] = []
    seen: set[str] = set()

    for row in rows:
        raw_bse = _cell(row, header_map, "bse_code")
        bse = str(raw_bse or "").strip()
        if bse.endswith(".0"):
            bse = bse[:-2]
        if not bse.isdigit():
            # Fall back: resolve NSE → BSE via master when possible.
            nse = str(_cell(row, header_map, "nse_symbol") or "").strip().upper()
            if not nse:
                skipped += 1
                continue
            from pms_platform.market_data.bse_scrip_universe import resolve_bse_code

            resolved = resolve_bse_code(nse_symbol=nse, company_name=str(_cell(row, header_map, "name") or ""))
            if not resolved or not str(resolved).isdigit():
                skipped += 1
                continue
            bse = str(resolved).strip()

        if bse in seen:
            continue
        seen.add(bse)
        codes.append(bse)

        resolution = resolver.resolve("BSE_CODE", bse, today)
        security_id = resolution.security_id if resolution.status == "RESOLVED" else None

        def d(field: str) -> Decimal | None:
            return _to_dec(_cell(row, header_map, field))

        mcap = d("market_cap_cr")
        pe = d("pe_ratio")
        earnings_yield = None
        if pe and pe != 0:
            earnings_yield = (Decimal("100") / pe).quantize(Decimal("0.0001"))

        upsert_valuation_snapshot(
            session,
            identifier_type="BSE_CODE",
            identifier=bse,
            security_id=security_id,
            as_of_date=today,
            last_price=None,
            market_cap_cr=mcap,
            pe_ratio=pe,
            industry_pe=d("industry_pe"),
            book_value_per_share=d("book_value_per_share"),
            price_to_book=d("price_to_book"),
            eps=d("eps"),
            dividend_yield=d("dividend_yield"),
            earnings_yield=earnings_yield,
            price_to_sales=d("price_to_sales"),
            peg_ratio=d("peg_ratio"),
            week_52_high=d("week_52_high"),
            week_52_low=d("week_52_low"),
            return_1d_pct=None,
            return_1m_pct=d("return_1m_pct"),
            return_3m_pct=d("return_3m_pct"),
            return_6m_pct=d("return_6m_pct"),
            return_1y_pct=d("return_1y_pct"),
            return_3y_pct=d("return_3y_pct"),
            provider="screener",
        )
        val_n += 1

        wrote_fund = _upsert_fund_snapshot(
            session,
            identifier_type="BSE_CODE",
            identifier=bse,
            security_id=security_id,
            period_end=period_end,
            sales=d("sales"),
            pat=d("pat"),
            opm=d("opm"),
            npm=d("npm"),
            sales_yoy_pct=d("sales_yoy_pct"),
            pat_yoy_pct=d("pat_yoy_pct"),
            sales_3y_cagr=d("sales_3y_cagr"),
            sales_5y_cagr=d("sales_5y_cagr"),
            pat_3y_cagr=d("pat_3y_cagr"),
            pat_5y_cagr=d("pat_5y_cagr"),
        )
        if wrote_fund or any(
            d(f) is not None
            for f in (
                "sales",
                "pat",
                "opm",
                "npm",
                "sales_yoy_pct",
                "pat_yoy_pct",
                "sales_3y_cagr",
                "sales_5y_cagr",
                "pat_3y_cagr",
                "pat_5y_cagr",
            )
        ):
            fund_n += 1

        roce = d("roce")
        roe = d("roe")
        roa = d("roa")
        de = d("debt_to_equity")
        ic = d("interest_coverage")
        cr = d("current_ratio")
        if any(v is not None for v in (roce, roe, roa, de, ic, cr)):
            fy, _ = _fiscal_from_period_end(period_end)
            upsert_annual_snapshot(
                session,
                identifier_type="BSE_CODE",
                identifier=bse,
                security_id=security_id,
                fiscal_year=fy,
                period_end_date=period_end,
                total_assets=None,
                total_equity=None,
                total_debt=None,
                cash_and_equivalents=None,
                finance_costs=None,
                current_assets=None,
                current_liabilities=None,
                roce=roce,
                roe=roe,
                roa=roa,
                debt_to_equity=de,
                interest_coverage=ic,
                current_ratio=cr,
                provider="screener",
            )
            ann_n += 1

        prom = d("promoter_holding_pct")
        pledge = d("pledged_pct")
        if prom is not None or pledge is not None:
            _upsert_promoter_snapshot(
                session,
                identifier_type="BSE_CODE",
                identifier=bse,
                security_id=security_id,
                quarter_end_date=period_end,
                promoter_holding_pct=prom,
                pledged_pct=pledge,
                prior_holding_pct=None,
                provider="screener",
            )
            prom_n += 1

    session.flush()
    return ScreenerExportImportResult(
        path=path,
        rows_read=len(rows),
        valuation_upserts=val_n,
        fund_upserts=fund_n,
        annual_upserts=ann_n,
        promoter_upserts=prom_n,
        skipped=skipped,
        bse_codes=tuple(codes),
    )
