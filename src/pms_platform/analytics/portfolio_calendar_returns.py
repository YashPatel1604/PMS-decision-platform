"""Official PMS performance series for portfolio / BSE period comparisons.

Ground truth (Research):

- ``Values.xlsx`` — monthly NAV for Portfolio + BSE SmallCap / MidCap / Sensex.
  YoY changes match ``CAGR_PMS.xlsx``. This is a *performance* path (not cash
  contributions into the live book).
- ``CAGR_PMS.xlsx`` — calendar-year returns (fallback / extension past Values).

Holdings ``PORT %`` / ``BSE %`` must use these series — never reconstructed
live AUM start→end (that overstates return when capital was added).
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

from pms_platform.config import settings

_ONE = Decimal("1")
_HUNDRED = Decimal("100")
_DAYS_PER_YEAR = Decimal("365")

PORTFOLIO_CALENDAR_REL = "portfolio/pms_calendar_returns.csv"
PORTFOLIO_NAV_REL = "portfolio/pms_nav_monthly.csv"


@dataclass(frozen=True)
class CalendarYearReturn:
    """One calendar-year total return (decimal fraction, e.g. 0.25 = +25%)."""

    year: int
    portfolio_return: Decimal
    bse_smallcap_return: Decimal | None = None
    source: str = "CAGR_PMS"


@dataclass(frozen=True)
class NavMonth:
    """One month-end performance snapshot from Values.xlsx."""

    as_of_date: date
    portfolio_nav: Decimal
    bse_smallcap: Decimal | None
    bse_midcap: Decimal | None
    sensex: Decimal | None
    monthly_return: Decimal | None
    source: str = "Values.xlsx"


def _candidate_paths(rel: str, external_dir: Path | None = None) -> list[Path]:
    candidates: list[Path] = []
    if external_dir is not None:
        candidates.append(Path(external_dir) / rel)
    candidates.append(Path(settings.external_data_dir) / rel)
    candidates.append(Path("/data/external") / rel)
    candidates.append(Path("/data/external_seed") / rel)
    candidates.append(Path("docker/market_data_seed") / rel)
    return candidates


def resolve_portfolio_calendar_path(external_dir: Path | None = None) -> Path | None:
    for path in _candidate_paths(PORTFOLIO_CALENDAR_REL, external_dir):
        if path.is_file():
            return path.resolve()
    return None


def resolve_portfolio_nav_path(external_dir: Path | None = None) -> Path | None:
    for path in _candidate_paths(PORTFOLIO_NAV_REL, external_dir):
        if path.is_file():
            return path.resolve()
    return None


def load_portfolio_calendar_returns(
    path: Path | None = None,
) -> dict[int, CalendarYearReturn]:
    resolved = path or resolve_portfolio_calendar_path()
    if resolved is None:
        return {}
    out: dict[int, CalendarYearReturn] = {}
    with resolved.open(newline="", encoding="utf-8") as handle:
        for raw in csv.DictReader(handle):
            year_raw = str(raw.get("year") or "").strip()
            if not year_raw:
                continue
            year = int(year_raw)
            if year < 100:
                year += 2000
            port = Decimal(str(raw["portfolio_return"]))
            bse_raw = raw.get("bse_smallcap_return")
            bse = Decimal(str(bse_raw)) if bse_raw not in (None, "") else None
            out[year] = CalendarYearReturn(
                year=year,
                portfolio_return=port,
                bse_smallcap_return=bse,
                source=str(raw.get("source") or "CAGR_PMS"),
            )
    return out


def load_portfolio_nav_months(path: Path | None = None) -> list[NavMonth]:
    resolved = path or resolve_portfolio_nav_path()
    if resolved is None:
        return []
    out: list[NavMonth] = []
    with resolved.open(newline="", encoding="utf-8") as handle:
        for raw in csv.DictReader(handle):
            if not raw.get("as_of_date") or not raw.get("portfolio_nav"):
                continue

            def _opt(key: str) -> Decimal | None:
                val = raw.get(key)
                if val in (None, ""):
                    return None
                return Decimal(str(val))

            out.append(
                NavMonth(
                    as_of_date=date.fromisoformat(str(raw["as_of_date"])[:10]),
                    portfolio_nav=Decimal(str(raw["portfolio_nav"])),
                    bse_smallcap=_opt("bse_smallcap"),
                    bse_midcap=_opt("bse_midcap"),
                    sensex=_opt("sensex"),
                    monthly_return=_opt("monthly_return"),
                    source=str(raw.get("source") or "Values.xlsx"),
                )
            )
    out.sort(key=lambda row: row.as_of_date)
    return out


@lru_cache(maxsize=1)
def _cached_calendar() -> tuple[tuple[int, CalendarYearReturn], ...]:
    return tuple(sorted(load_portfolio_calendar_returns().items()))


@lru_cache(maxsize=1)
def _cached_nav() -> tuple[NavMonth, ...]:
    return tuple(load_portfolio_nav_months())


def clear_portfolio_calendar_cache() -> None:
    _cached_calendar.cache_clear()
    _cached_nav.cache_clear()


def _nav_on_or_before(as_of: date, months: list[NavMonth]) -> NavMonth | None:
    chosen: NavMonth | None = None
    for row in months:
        if row.as_of_date <= as_of:
            chosen = row
        else:
            break
    return chosen


def nav_period_return_pct(
    start_date: date,
    end_date: date,
    *,
    field: str = "portfolio_nav",
    months: list[NavMonth] | None = None,
) -> tuple[Decimal, str] | None:
    """Return (total_return_pct, methodology) from Values NAV levels."""
    series = months if months is not None else list(_cached_nav())
    if not series:
        return None
    start = _nav_on_or_before(start_date, series)
    end = _nav_on_or_before(end_date, series)
    if start is None or end is None:
        return None
    start_level = getattr(start, field)
    end_level = getattr(end, field)
    if start_level is None or end_level is None or start_level <= 0:
        return None
    total = ((end_level / start_level) - _ONE) * _HUNDRED
    method = f"VALUES_NAV:{field}:{start.as_of_date.isoformat()}->{end.as_of_date.isoformat()}"
    return total, method


def linked_calendar_return_pct(
    start_date: date,
    end_date: date,
    *,
    use_bse_smallcap: bool = False,
    calendar: dict[int, CalendarYearReturn] | None = None,
) -> Decimal | None:
    """Link CAGR calendar-year returns across a window (partial years scaled)."""
    if end_date < start_date:
        return None
    years = calendar if calendar is not None else dict(_cached_calendar())
    if not years:
        return None

    growth = _ONE
    year = start_date.year
    while year <= end_date.year:
        row = years.get(year)
        if row is None:
            return None
        rate = row.bse_smallcap_return if use_bse_smallcap else row.portfolio_return
        if rate is None:
            return None
        year_start = date(year, 1, 1)
        year_end = date(year, 12, 31)
        window_start = max(start_date, year_start)
        window_end = min(end_date, year_end)
        days = (window_end - window_start).days + 1
        if days <= 0:
            year += 1
            continue
        if window_start == year_start and window_end == year_end:
            growth *= _ONE + rate
        else:
            fraction = Decimal(days) / _DAYS_PER_YEAR
            growth *= (_ONE + rate) ** fraction
        year += 1
    return (growth - _ONE) * _HUNDRED


def linked_portfolio_return_pct(
    start_date: date,
    end_date: date,
    *,
    calendar: dict[int, CalendarYearReturn] | None = None,
) -> Decimal | None:
    """Prefer Values NAV when the window fits; else CAGR calendar TWR.

    If ``calendar`` is passed explicitly (tests), skip NAV and use that map.
    """
    if calendar is not None:
        return linked_calendar_return_pct(start_date, end_date, calendar=calendar)

    months = list(_cached_nav())
    last = months[-1].as_of_date if months else None
    if last is not None and end_date <= last:
        nav = nav_period_return_pct(
            start_date, end_date, field="portfolio_nav", months=months
        )
        if nav is not None:
            return nav[0]
    return linked_calendar_return_pct(start_date, end_date)


def linked_bse_smallcap_return_pct(
    start_date: date,
    end_date: date,
) -> Decimal | None:
    """BSE SmallCap from Values NAV when possible; else CAGR calendar."""
    months = list(_cached_nav())
    last = months[-1].as_of_date if months else None
    if last is not None and end_date <= last:
        nav = nav_period_return_pct(start_date, end_date, field="bse_smallcap", months=months)
        if nav is not None:
            return nav[0]
    return linked_calendar_return_pct(start_date, end_date, use_bse_smallcap=True)


# Back-compat alias used by older call sites / tests
def linked_portfolio_return_pct_calendar_only(
    start_date: date,
    end_date: date,
    *,
    calendar: dict[int, CalendarYearReturn] | None = None,
) -> Decimal | None:
    return linked_calendar_return_pct(start_date, end_date, calendar=calendar)


def write_portfolio_calendar_csv(path: Path, rows: list[CalendarYearReturn]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["year", "portfolio_return", "bse_smallcap_return", "source"],
        )
        writer.writeheader()
        for row in sorted(rows, key=lambda item: item.year):
            writer.writerow(
                {
                    "year": row.year,
                    "portfolio_return": f"{row.portfolio_return:f}",
                    "bse_smallcap_return": (
                        f"{row.bse_smallcap_return:f}"
                        if row.bse_smallcap_return is not None
                        else ""
                    ),
                    "source": row.source,
                }
            )


def extract_calendar_from_cagr_workbook(path: Path) -> list[CalendarYearReturn]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["From Start"] if "From Start" in workbook.sheetnames else workbook.active
    rows: list[CalendarYearReturn] = []
    for raw in sheet.iter_rows(values_only=True):
        if not raw or raw[0] is None or raw[1] is None:
            continue
        label = str(raw[0]).strip()
        if not label.upper().startswith("CY"):
            continue
        year = int(label[2:])
        if year < 100:
            year += 2000
        port = Decimal(str(raw[1]))
        bse = Decimal(str(raw[2])) if raw[2] is not None else None
        rows.append(
            CalendarYearReturn(
                year=year,
                portfolio_return=port,
                bse_smallcap_return=bse,
                source="CAGR_PMS",
            )
        )
    workbook.close()
    return rows
