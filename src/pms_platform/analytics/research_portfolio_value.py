"""Authoritative portfolio market values from Research/Portfolio workbooks.

Lookup order for a target date (History always beats Model Portfolio when available):
1. ``History/PMS_ClientPortfolio_DDMMYY`` Model sheet ``Total_Value`` (exact date)
2. Closest History observation on/before the target date
3. Exact ``Portfolio_YYYY.xlsx`` / Portfolio Yearly sheet total (``MODEL_PORTFOLIO``)
4. Exact ``Values.xlsx`` Date → Portfolio Value
5. Closest Model Portfolio observation on/before
6. Closest Values observation on/before

``MODEL_PORTFOLIO`` usually means Docker/app could not read ``Research/Portfolio/History``
(or no History file exists for that window) and fell back to yearly snapshot books.

Reconstructed ledger valuation remains a separate cross-check in open holdings.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import lru_cache
from pathlib import Path

import openpyxl

from pms_platform.ingestion.snapshots import parse_snapshot_sheet_date
from pms_platform.research_paths import (
    portfolio_snapshot_source_dirs,
    research_portfolio_history_dir,
    research_values_workbook,
)

_ZERO = Decimal("0")
_SKIP_LABELS = frozenset({"stock", "total", "total_value", "cash", "model", "qnty"})


@dataclass(frozen=True)
class ResearchPortfolioValue:
    """One Research-sourced portfolio total."""

    as_of_date: date
    value: Decimal
    source: str
    source_file: str
    observation_date: date


def _to_decimal(raw: object) -> Decimal | None:
    if raw is None:
        return None
    try:
        return Decimal(str(raw))
    except Exception:  # noqa: BLE001
        return None


def _history_token(as_of_date: date) -> str:
    return f"{as_of_date.day:02d}{as_of_date.month:02d}{as_of_date.year % 100:02d}"


def _parse_history_filename_date(path: Path) -> date | None:
    match = re.search(r"PMS_ClientPortfolio_(\d{6})\.", path.name, flags=re.IGNORECASE)
    if match is None:
        return None
    token = match.group(1)
    day, month, year = int(token[:2]), int(token[2:4]), 2000 + int(token[4:6])
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _model_total_from_rows(rows: list[tuple[object, ...]]) -> Decimal | None:
    total_label: Decimal | None = None
    summed = _ZERO
    counted = 0
    for row in rows:
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        if not label:
            continue
        lower = label.lower().replace(" ", "_")
        amount = _to_decimal(row[3] if len(row) > 3 else None)
        if lower in {"total_value", "total"} and amount is not None:
            total_label = amount
            continue
        if lower in _SKIP_LABELS:
            continue
        if amount is None:
            continue
        summed += amount
        counted += 1
    if total_label is not None:
        return total_label
    return summed if counted else None


def _history_model_total(path: Path) -> Decimal | None:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                if "Model" not in workbook.sheetnames:
                    return None
                rows = [tuple(row) for row in workbook["Model"].iter_rows(values_only=True)]
            finally:
                workbook.close()
            return _model_total_from_rows(rows)
        if suffix == ".xls":
            import xlrd

            book = xlrd.open_workbook(path)
            if "Model" not in book.sheet_names():
                return None
            sheet = book.sheet_by_name("Model")
            rows = [tuple(sheet.row_values(r)) for r in range(sheet.nrows)]
            return _model_total_from_rows(rows)
    except Exception:  # noqa: BLE001
        return None
    return None


def _find_exact_history_file(as_of_date: date) -> Path | None:
    history_dir = research_portfolio_history_dir()
    if history_dir is None:
        return None
    token = _history_token(as_of_date)
    for path in sorted(history_dir.glob(f"PMS_ClientPortfolio_{token}.*")):
        if path.suffix.lower() in {".xlsx", ".xls"}:
            return path
    return None


def _load_history_dates() -> dict[date, Path]:
    history_dir = research_portfolio_history_dir()
    if history_dir is None:
        return {}
    index: dict[date, Path] = {}
    for path in history_dir.glob("PMS_ClientPortfolio_*"):
        if path.suffix.lower() not in {".xlsx", ".xls"}:
            continue
        as_of = _parse_history_filename_date(path)
        if as_of is None:
            continue
        # Prefer .xlsx over .xls when both exist.
        existing = index.get(as_of)
        if existing is None or (
            existing.suffix.lower() == ".xls" and path.suffix.lower() == ".xlsx"
        ):
            index[as_of] = path
    return index


def _model_portfolio_sheet_total(path: Path, sheet_name: str) -> Decimal | None:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
        finally:
            workbook.close()
    except Exception:  # noqa: BLE001
        return None
    total = _ZERO
    counted = 0
    for idx, row in enumerate(rows):
        if idx < 4:
            continue
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        if not label or label.lower() in {"total", "cash", "stock"}:
            continue
        amount = _to_decimal(row[3] if len(row) > 3 else None)
        if amount is None:
            continue
        total += amount
        counted += 1
    return total if counted else None


@lru_cache(maxsize=1)
def _model_portfolio_index() -> dict[date, tuple[Decimal, Path, str]]:
    index: dict[date, tuple[Decimal, Path, str]] = {}
    for directory in portfolio_snapshot_source_dirs():
        for path in sorted(directory.glob("Portfolio_*.xlsx")):
            year_match = re.search(r"Portfolio_(\d{4})", path.name)
            if year_match is None:
                continue
            workbook_year = int(year_match.group(1))
            try:
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            except Exception:  # noqa: BLE001
                continue
            try:
                for sheet_name in workbook.sheetnames:
                    if sheet_name.strip().lower() in {"liquid", "sheet1", "sheet2", "sheet3"}:
                        continue
                    snapshot_date = parse_snapshot_sheet_date(sheet_name, workbook_year)
                    if snapshot_date is None:
                        header = next(
                            workbook[sheet_name].iter_rows(min_row=1, max_row=1, values_only=True),
                            None,
                        )
                        if header and len(header) > 1 and isinstance(header[1], datetime):
                            snapshot_date = header[1].date()
                    if snapshot_date is None or snapshot_date in index:
                        continue
                    total = _model_portfolio_sheet_total(path, sheet_name)
                    if total is None:
                        continue
                    index[snapshot_date] = (total, path, sheet_name)
            finally:
                workbook.close()
    return index


@lru_cache(maxsize=1)
def _values_series() -> dict[date, Decimal]:
    path = research_values_workbook()
    if path is None:
        return {}
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [tuple(row) for row in workbook["Sheet1"].iter_rows(values_only=True)]
        finally:
            workbook.close()
    except Exception:  # noqa: BLE001
        return {}
    series: dict[date, Decimal] = {}
    for row in rows[1:]:
        if not row:
            continue
        raw_date, raw_value = row[0], row[1] if len(row) > 1 else None
        if isinstance(raw_date, datetime):
            day = raw_date.date()
        elif isinstance(raw_date, date):
            day = raw_date
        else:
            continue
        value = _to_decimal(raw_value)
        if value is None:
            continue
        series[day] = value
    return series


@lru_cache(maxsize=1)
def _history_date_index() -> dict[date, Path]:
    return _load_history_dates()


def clear_research_portfolio_value_cache() -> None:
    """Drop cached Research workbook indexes (tests / after sync)."""
    _model_portfolio_index.cache_clear()
    _values_series.cache_clear()
    _history_date_index.cache_clear()


def latest_research_book_date() -> date | None:
    """Latest Excel book date for Current Holdings (History preferred).

    Current Holdings must not run past the last Research workbook
    observation. Prefer ``History/PMS_ClientPortfolio_*``; fall back to
    Model Portfolio / Values only when History is unavailable.
    """
    history = _history_date_index()
    if history:
        return max(history)
    models = _model_portfolio_index()
    values = _values_series()
    candidates = list(models) + list(values)
    return max(candidates) if candidates else None


def _on_or_before(mapping_dates: list[date], target: date) -> date | None:
    prior = [day for day in mapping_dates if day <= target]
    return max(prior) if prior else None


def lookup_research_portfolio_value(as_of_date: date) -> ResearchPortfolioValue | None:
    """Return the best Research portfolio total for ``as_of_date``."""
    exact_history = _find_exact_history_file(as_of_date)
    if exact_history is not None:
        total = _history_model_total(exact_history)
        if total is not None:
            return ResearchPortfolioValue(
                as_of_date=as_of_date,
                value=total,
                source="HISTORY",
                source_file=str(exact_history),
                observation_date=as_of_date,
            )

    # Prefer any History observation on/before target before yearly Model books.
    history_index = _history_date_index()
    history_day = _on_or_before(list(history_index), as_of_date)
    if history_day is not None:
        path = history_index[history_day]
        total = _history_model_total(path)
        if total is not None:
            return ResearchPortfolioValue(
                as_of_date=as_of_date,
                value=total,
                source="HISTORY",
                source_file=str(path),
                observation_date=history_day,
            )

    models = _model_portfolio_index()
    if as_of_date in models:
        value, path, sheet = models[as_of_date]
        return ResearchPortfolioValue(
            as_of_date=as_of_date,
            value=value,
            source="MODEL_PORTFOLIO",
            source_file=f"{path}#{sheet}",
            observation_date=as_of_date,
        )

    values = _values_series()
    if as_of_date in values:
        return ResearchPortfolioValue(
            as_of_date=as_of_date,
            value=values[as_of_date],
            source="VALUES",
            source_file=str(research_values_workbook()),
            observation_date=as_of_date,
        )

    model_day = _on_or_before(list(models), as_of_date)
    if model_day is not None:
        value, path, sheet = models[model_day]
        return ResearchPortfolioValue(
            as_of_date=as_of_date,
            value=value,
            source="MODEL_PORTFOLIO",
            source_file=f"{path}#{sheet}",
            observation_date=model_day,
        )

    values_day = _on_or_before(list(values), as_of_date)
    if values_day is not None:
        return ResearchPortfolioValue(
            as_of_date=as_of_date,
            value=values[values_day],
            source="VALUES",
            source_file=str(research_values_workbook()),
            observation_date=values_day,
        )

    return None
