"""DailyEditFiles Charts + SCA_LLP stay in lockstep with committed bhav."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from pms_platform.market_data.daily_edit_bhav import (
    apply_bhav_csv_to_daily_edit_files,
    client_portfolio_daily_edit_path,
    pivot_workbook_daily_edit_path,
)
from pms_platform.market_data.client_portfolio_parse import client_portfolio_workbook_path
from pms_platform.market_data.pivot_seed import default_pivot_workbook


def test_apply_bhav_updates_charts_and_sca(tmp_path: Path) -> None:
    csv_path = tmp_path / "bhav.csv"
    csv_path.write_text(
        "TradDt,TckrSymb,SctySrs,OpnPric,HghPric,LwPric,ClsPric\n"
        "2026-08-21,ASHAPURMIN,EQ,100,110,90,105\n",
        encoding="utf-8",
    )

    charts = tmp_path / "Charts - Copy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BhavCopy_NSE_CM"
    ws.append(["SYMBOL+Series", "TradDt", "TckrSymb", "SctySrs", "ClsPric"])
    ws.append(["=CONCATENATE(C2,D2)", "2026-08-20", "OLD", "EQ", 1])
    wb.create_sheet("Range ")
    wb.save(charts)
    wb.close()

    sca = tmp_path / "SCA_LLP Stock Holding - Copy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "cmbhavcopy"
    ws.append(["TradDt", "TckrSymb", "SctySrs", "ClsPric"])
    qty = wb.create_sheet("Quantity")
    qty.append(
        [
            "Symbol",
            "Code",
            "Percentage",
            "Total Quantity",
            "PRICE",
            "VALUE",
            "old",
            "RAMPRASATH REDDY QTYN",
            "BLOCKED ACCOUNT",
        ]
    )
    qty.append(["ASHAPURMIN", 1, 0.1, 10, 1, 10, "=D2-H2", 3, "=H2*E2"])
    stocks = wb.create_sheet("Stocks")
    stocks.append(["SYMBOL", "Stocks", None])
    stocks.append(["ASHAPURMIN", "ASHAPURMIN", 1])
    wb.save(sca)
    wb.close()

    result = apply_bhav_csv_to_daily_edit_files(csv_path, folder=tmp_path)
    assert result.get("error") is None
    assert result["charts_rows"] == 1
    assert result["sca_bhav_rows"] == 1
    assert result["sca_qty_priced"] == 1
    assert result["sca_stocks_priced"] == 1

    charts_wb = load_workbook(charts, data_only=False)
    row = next(charts_wb["BhavCopy_NSE_CM"].iter_rows(min_row=2, max_row=2, values_only=True))
    assert row[2] == "ASHAPURMIN"
    assert str(row[0]).startswith("=CONCATENATE")
    charts_wb.close()

    sca_wb = load_workbook(sca, data_only=False)
    qty_ws = sca_wb["Quantity"]
    assert qty_ws["E2"].value == 105
    assert qty_ws["F2"].value == 1050
    assert qty_ws["G1"].value == "21.08.2026"
    assert qty_ws["G2"].value == "=D2-H2"
    assert qty_ws["H2"].value == 3
    assert qty_ws["I2"].value == "=H2*E2"
    sca_wb.close()


def test_finds_client_and_pivot_workbooks_in_daily_edit(tmp_path: Path, monkeypatch) -> None:
    client = tmp_path / "PMS_ClientPortfolio copy.xlsx"
    pivot = tmp_path / "PivotPointsStrategy_New - Copy.xlsx"
    client.write_bytes(b"PK")
    pivot.write_bytes(b"PK")
    monkeypatch.setattr(
        "pms_platform.market_data.daily_edit_bhav.daily_edit_dir",
        lambda: tmp_path,
    )
    assert client_portfolio_daily_edit_path().name == client.name
    assert pivot_workbook_daily_edit_path().name == pivot.name
    assert client_portfolio_workbook_path() == client
    assert default_pivot_workbook() == pivot
