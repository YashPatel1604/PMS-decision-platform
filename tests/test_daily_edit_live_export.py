"""Live export patches SCA totals from software holdings."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from pms_platform.market_data.daily_edit_live_export import patch_live_workbook


def _sca_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quantity"
    ws.append(["Symbol", "Code", "Percentage", "Total Quantity", "PRICE", "VALUE", "date"])
    ws.append(["ASHAPURMIN", None, None, 100, 10, "=E2*D2"])
    ws.append([None, None, None, None, None, "=SUM(F2:F2)"])
    ws.append(["Balance with Bank", None, None, None, None, 1000])
    ws.append([None, None, None, None, None, "=F3+F4"])
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_patch_sca_writes_live_price_value_and_totals(monkeypatch) -> None:
    def _fake_dash(session, *, book="sca", **_kwargs):
        assert book == "sca"
        return {
            "as_of": "2026-09-10",
            "bank_balance": 5000.0,
            "total_value": 2500.0,
            "portfolio_total": 7500.0,
            "holdings": [
                {"symbol": "ASHAPURMIN", "qty": 100.0, "price": 25.0, "value": 2500.0},
            ],
        }

    monkeypatch.setattr(
        "pms_platform.market_data.client_portfolio_dashboard.build_client_portfolio_dashboard",
        _fake_dash,
    )
    out = patch_live_workbook(object(), "sca_llp", _sca_template())  # type: ignore[arg-type]
    ws = load_workbook(BytesIO(out))["Quantity"]
    assert ws.cell(2, 5).value == 25.0
    assert ws.cell(2, 6).value == 2500.0
    assert ws.cell(3, 6).value == 2500.0  # equity total
    assert ws.cell(4, 6).value == 5000.0  # bank
    assert ws.cell(5, 6).value == 7500.0  # portfolio total
    assert ws.cell(1, 7).value == "10.09.2026"
