"""DailyEdit upload and reimport tests."""

from __future__ import annotations

import io

import openpyxl
import pytest

from pms_platform.auth.passwords import hash_password
from pms_platform.ingestion.daily_edit_sync import (
    daily_edit_status,
    reimport_daily_edit,
    upload_daily_edit,
)
from pms_platform.models.user import User
from pms_platform.storage.adapter import LocalFilesystemStorage


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Range"
    ws["A1"] = "Name"
    ws["B1"] = "High"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def storage(tmp_path, monkeypatch):
    adapter = LocalFilesystemStorage(tmp_path / "private")
    monkeypatch.setattr("pms_platform.storage.get_storage", lambda: adapter)
    return adapter


@pytest.fixture
def daily_edit_dir(tmp_path, monkeypatch):
    edit = tmp_path / "daily_edit"
    edit.mkdir()
    monkeypatch.setattr("pms_platform.config.settings.daily_edit_dir", edit)
    return edit


def _admin(session) -> User:
    user = User(
        email="admin@t.com",
        password_hash=hash_password("x"),
        display_name="admin",
        role="admin",
    )
    session.add(user)
    session.flush()
    return user


def test_upload_materializes_and_dedupes(session, storage, daily_edit_dir) -> None:
    admin = _admin(session)
    data = _xlsx_bytes()
    first = upload_daily_edit(
        session,
        category="charts",
        filename="Charts_MCAP.xlsx",
        data=data,
        uploaded_by=admin.user_id,
        storage=storage,
    )
    assert first.deduplicated is False
    assert (daily_edit_dir / "Charts.xlsx").is_file()

    second = upload_daily_edit(
        session,
        category="charts",
        filename="Charts_MCAP.xlsx",
        data=data,
        uploaded_by=admin.user_id,
        storage=storage,
    )
    assert second.deduplicated is True

    status = daily_edit_status(session)
    assert status["categories"]["charts"]["uploaded"] is True
    assert status["categories"]["charts"]["on_disk"] is True


def test_reimport_charts_dry_run(session, storage, daily_edit_dir, monkeypatch) -> None:
    admin = _admin(session)
    data = _xlsx_bytes()
    upload_daily_edit(
        session,
        category="charts",
        filename="Charts.xlsx",
        data=data,
        uploaded_by=admin.user_id,
        storage=storage,
    )
    monkeypatch.setattr(
        "pms_platform.ingestion.daily_edit_sync.parse_charts_range",
        lambda _path: [
            {"excel_row": 2, "symbol": "TCS", "name": "TCS", "section": "holdings"},
        ],
    )
    preview = reimport_daily_edit(session, category="charts", dry_run=True)
    assert preview["dry_run"] is True
    assert preview["added"] == [2]


def test_status_marks_sca_and_pivot_reimportable(session, daily_edit_dir) -> None:
    status = daily_edit_status(session)
    assert status["categories"]["sca_llp"]["reimport_supported"] is True
    assert status["categories"]["pivot_points"]["reimport_supported"] is True


def test_reimport_sca_uses_sca_book(session, storage, daily_edit_dir, monkeypatch) -> None:
    admin = _admin(session)
    data = _xlsx_bytes()
    upload_daily_edit(
        session,
        category="sca_llp",
        filename="SCA.xlsx",
        data=data,
        uploaded_by=admin.user_id,
        storage=storage,
    )

    class _Pos:
        symbol = "RELIANCE"
        qty = 10
        mcap_factor = None
        index_label = None

    class _Book:
        model = [_Pos()]

    monkeypatch.setattr(
        "pms_platform.market_data.client_portfolio_parse.parse_client_portfolio_workbook",
        lambda _path: _Book(),
    )
    calls: list[str] = []

    def _fake_reimport(session, holdings, *, book="client", update_qty=False, dry_run=False):
        calls.append(book)
        return {"added": ["RELIANCE"], "removed": [], "metadata_updated": [], "qty_updated": []}

    monkeypatch.setattr(
        "pms_platform.ingestion.daily_edit_sync.reimport_client_positions",
        _fake_reimport,
    )
    preview = reimport_daily_edit(session, category="sca_llp", dry_run=True)
    assert preview["added"] == ["RELIANCE"]
    assert calls == ["sca"]


def test_reimport_pivot_dry_run(session, storage, daily_edit_dir, monkeypatch) -> None:
    admin = _admin(session)
    data = _xlsx_bytes()
    upload_daily_edit(
        session,
        category="pivot_points",
        filename="Pivot.xlsx",
        data=data,
        uploaded_by=admin.user_id,
        storage=storage,
    )
    monkeypatch.setattr(
        "pms_platform.market_data.pivot_seed.reimport_pivot_portfolio",
        lambda *_a, **_k: {
            "added": ["TCS"],
            "removed": [],
            "metadata_updated": [],
            "qty_updated": [],
            "portfolio_symbols": 1,
        },
    )
    preview = reimport_daily_edit(session, category="pivot_points", dry_run=True)
    assert preview["added"] == ["TCS"]
    assert preview["portfolio_symbols"] == 1
