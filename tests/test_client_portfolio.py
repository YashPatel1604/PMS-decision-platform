"""Client Portfolio workbook parse + bhav join."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from pms_platform.market_data.client_portfolio_dashboard import (
    build_client_portfolio_dashboard,
)
from pms_platform.market_data.client_portfolio_parse import (
    clear_client_portfolio_cache,
    parse_client_portfolio_workbook,
    write_bse_smallcap_year,
)
from pms_platform.market_data.nse_bhav_store import sync_bhav_file, upsert_portfolio_symbols
from pms_platform.market_data.pivot_derived import floor_pivot_levels

FIXTURES = Path(__file__).parent / "fixtures" / "pivot"
BOOK = FIXTURES / "PMS_ClientPortfolio_fixture.xlsx"


def test_parse_client_portfolio_model_and_stocks() -> None:
    book = parse_client_portfolio_workbook(BOOK)
    assert book.excel_total_value == Decimal("44305")
    by_sym = {p.symbol: p for p in book.model}
    assert set(by_sym) == {"RELIANCE", "INFY", "NOSUCH"}
    assert by_sym["RELIANCE"].qty == Decimal(10)
    assert book.stocks_qty["RELIANCE"] == Decimal(100)


def test_client_portfolio_dashboard_joins_bhav(session, tmp_path, monkeypatch) -> None:
    clear_client_portfolio_cache()
    monkeypatch.setattr(
        "pms_platform.market_data.nse_bhav_store.settings.upload_dir",
        str(tmp_path),
    )
    monkeypatch.setattr(
        "pms_platform.market_data.client_portfolio_dashboard.load_client_portfolio_book",
        lambda path=None: parse_client_portfolio_workbook(BOOK),
    )
    sync_bhav_file(session, FIXTURES / "bhav_2026-08-19.csv")
    upsert_portfolio_symbols(session, [{"symbol": "RELIANCE", "portfolio_a": True}])
    session.commit()

    dash = build_client_portfolio_dashboard(session, as_of=date(2026, 8, 19))
    assert dash["error"] is None
    assert dash["as_of"] == "2026-08-19"
    by_sym = {r["symbol"]: r for r in dash["holdings"]}
    rel = by_sym["RELIANCE"]
    assert rel["missing_bhav"] is False
    assert rel["close"] == 1425.0
    assert rel["qty"] == 10.0
    assert abs(rel["bhav_value"] - 14250.0) < 1e-6
    assert rel["qty_mismatch"] is True  # Model 10 vs Stocks 100
    levels = floor_pivot_levels(Decimal("1430"), Decimal("1400"), Decimal("1425"))
    assert abs(rel["pivot"]["pp"] - float(levels["pp"])) < 1e-9
    assert abs(rel["pivot"]["s4_03"] - float(levels["s4_03"])) < 1e-9
    assert by_sym["NOSUCH"]["missing_bhav"] is True
    assert "NOSUCH" in dash["missing_symbols"]


def test_sca_quantity_ramprasath_columns(tmp_path: Path) -> None:
    """Quantity!H is static; G=D−H; I=H×price (same as Excel)."""
    from openpyxl import Workbook

    from pms_platform.market_data.client_portfolio_parse import (
        clear_client_portfolio_cache,
        parse_client_portfolio_workbook,
    )

    path = tmp_path / "SCA_LLP Stock Holding - Copy.xlsx"
    wb = Workbook()
    wb.active.title = "cmbhavcopy"
    qty = wb.create_sheet("Quantity")
    qty.append(
        [
            "Symbol",
            "Code",
            "Percentage",
            "Total Quantity",
            "PRICE",
            "VALUE",
            "21.08.2026",
            "RAMPRASATH REDDY QTYN",
            "BLOCKED ACCOUNT",
        ]
    )
    qty.append(["ASHAPURMIN", 1, None, 100, 10, 1000, "=D2-H2", 15, "=H2*E2"])
    qty.append(["WELENT", 2, None, "=267283-62869-123080-64967", 50, None, None, 16367, None])
    qty.append(["Balance with Bank", None, None, None, None, 1_500_000])
    wb.save(path)
    wb.close()

    clear_client_portfolio_cache()
    book = parse_client_portfolio_workbook(path)
    by_sym = {p.symbol: p for p in book.model}
    assert by_sym["ASHAPURMIN"].qty == Decimal(100)
    assert by_sym["ASHAPURMIN"].ramprasath_qty == Decimal(15)
    assert by_sym["WELENT"].qty == Decimal(16367)  # 267283-62869-123080-64967
    assert by_sym["WELENT"].ramprasath_qty == Decimal(16367)

    from pms_platform.market_data.client_portfolio_dashboard import _fill_ramprasath

    row = {"price": 105.0, "ramprasath_qty": None, "ex_ramprasath_qty": None, "blocked_value": None}
    _fill_ramprasath(row, by_sym["ASHAPURMIN"])
    assert row["ramprasath_qty"] == 15.0
    assert row["ex_ramprasath_qty"] == 85.0
    assert row["blocked_value"] == 1575.0
    assert book.bank_balance == Decimal("1500000")
    assert "BALANCE WITH BANK" not in by_sym


def test_write_sca_bank_balance_updates_quantity_f(tmp_path: Path) -> None:
    from decimal import Decimal

    from openpyxl import Workbook, load_workbook

    from pms_platform.market_data.client_portfolio_parse import (
        clear_client_portfolio_cache,
        parse_client_portfolio_workbook,
    )
    from pms_platform.market_data.daily_edit_bhav import write_sca_bank_balance

    path = tmp_path / "SCA_LLP Stock Holding - Copy.xlsx"
    wb = Workbook()
    wb.active.title = "cmbhavcopy"
    qty = wb.create_sheet("Quantity")
    qty.append(["Symbol", "Code", "Percentage", "Total Quantity", "PRICE", "VALUE"])
    qty.append(["ASHAPURMIN", 1, None, 10, 100, 1000])
    qty.append(["Balance with Bank", None, None, None, None, 50])
    wb.save(path)
    wb.close()

    written = write_sca_bank_balance(Decimal("123456.78"), folder=tmp_path)
    assert written == 123456.78
    clear_client_portfolio_cache()
    book = parse_client_portfolio_workbook(path)
    assert book.bank_balance == Decimal("123456.78")
    wb2 = load_workbook(path, data_only=True)
    assert wb2["Quantity"]["F3"].value == 123456.78
    wb2.close()


def test_write_bse_smallcap_year_updates_workbook(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "PMS_ClientPortfolio_New.xlsx"
    wb = Workbook()
    model = wb.active
    model.title = "Model"
    model.cell(1, 12, "Portfolio")
    model.cell(1, 18, "BSESmallCap")
    model.cell(2, 12, 2025)
    model.cell(2, 13, 1000)
    model.cell(2, 14, 1100)
    model.cell(2, 18, 2025)
    model.cell(2, 19, 5000)
    model.cell(2, 20, 5500)
    model.cell(3, 12, 2026)
    model.cell(3, 13, 1100)
    model.cell(3, 14, 1200)
    model.cell(3, 18, 2026)
    model.cell(3, 19, 5500)
    model.cell(3, 20, 6000)
    model.cell(17, 11, "BSEMidCap")
    wb.create_sheet("Stocks")
    wb.save(path)
    wb.close()

    result = write_bse_smallcap_year(
        year=2026,
        start=Decimal("5600"),
        end=Decimal("6100"),
        path=path,
    )
    assert result == {"year": 2026, "start": 5600.0, "end": 6100.0}
    wb2 = load_workbook(path, data_only=True)
    assert wb2["Model"].cell(3, 19).value == 5600.0
    assert wb2["Model"].cell(3, 20).value == 6100.0
    wb2.close()


def test_mcap_firm_from_bhav_not_excel_cache(tmp_path: Path, session, monkeypatch) -> None:
    """Share factor stays in Excel; Mcap/%Firm revalue off as-of bhav close."""
    from openpyxl import Workbook

    from pms_platform.market_data.client_portfolio_dashboard import (
        build_client_portfolio_dashboard,
    )
    from pms_platform.market_data.nse_bhav_store import sync_bhav_file

    path = tmp_path / "PMS_ClientPortfolio copy.xlsx"
    wb = Workbook()
    model = wb.active
    model.title = "Model"
    model.append(
        ["Model", "Qnty", "Price", "Value", "Percent", "Index", "Mcap", "Date", "%Firm"]
    )
    model.append([None] * 9)
    model.append(
        [
            "RELIANCE",
            10,
            "=Stocks!C3",
            None,
            None,
            "-",
            "=(2/1)*C3",
            "Nov25",
            "=(Stocks!E3*C3)/(G3*100000)",
        ]
    )
    stocks = wb.create_sheet("Stocks")
    stocks.append(["SYMBOL", "Stocks", None, None, "Quantity", "Value"])
    stocks.append([None] * 6)
    stocks.append(["RELIANCE", "RELIANCE", 100.0, None, 50_000, None])  # stale Excel price
    wb.save(path)
    wb.close()

    clear_client_portfolio_cache()
    monkeypatch.setattr(
        "pms_platform.market_data.nse_bhav_store.settings.upload_dir",
        str(tmp_path / "uploads"),
    )
    monkeypatch.setattr(
        "pms_platform.market_data.client_portfolio_dashboard.load_client_portfolio_book",
        lambda _path=None: parse_client_portfolio_workbook(path),
    )
    sync_bhav_file(session, FIXTURES / "bhav_2026-08-19.csv")
    dash = build_client_portfolio_dashboard(session, as_of=date(2026, 8, 19))
    rel = next(r for r in dash["holdings"] if r["symbol"] == "RELIANCE")
    # Fixture RELIANCE close = 1425; factor = 2 → mcap = 2850
    assert rel["close"] == 1425.0
    assert abs(rel["mcap"] - 2850.0) < 1e-6
    assert abs(rel["firm_pct"] - (50000 * 1425) / (2850 * 100_000)) < 1e-9


def test_mcap_factor_parsed_when_excel_cache_missing(tmp_path: Path) -> None:
    """openpyxl save wipes formula caches — share factor still comes from the formula text."""
    from openpyxl import Workbook

    path = tmp_path / "PMS_ClientPortfolio copy.xlsx"
    wb = Workbook()
    model = wb.active
    model.title = "Model"
    model.append(
        ["Model", "Qnty", "Price", "Value", "Percent", "Index", "Mcap", "Date", "%Firm"]
    )
    model.append([None] * 9)
    model.append(
        [
            "ASHAPURMIN",
            3670,
            "=Stocks!C3",
            None,
            None,
            "-",
            "=(19.11/2)*C3",
            "Nov25",
            "=(Stocks!E3*C3)/(G3*100000)",
        ]
    )
    stocks = wb.create_sheet("Stocks")
    stocks.append(["SYMBOL", "Stocks", None, None, "Quantity", "Value"])
    stocks.append([None] * 6)
    stocks.append(["ASHAPURMIN", "ASHAPURMIN", 585.8, None, 499364, None])
    wb.save(path)
    wb.close()

    book = parse_client_portfolio_workbook(path)
    pos = book.model[0]
    assert pos.symbol == "ASHAPURMIN"
    assert pos.mcap_factor is not None
    assert abs(pos.mcap_factor - Decimal("19.11") / 2) < Decimal("0.0001")

