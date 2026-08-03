"""Tests for Final Master prompt parser and Excel apply."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pytest
from pms_platform.masters.apply import apply_edits
from pms_platform.masters.parser import parse_prompt
from pms_platform.masters.paths import MasterKind, resolve_master_path
from pms_platform.masters.preview import preview_workbook


def test_parse_buy_sell_and_update() -> None:
    result = parse_prompt(
        "\n".join(
            [
                "BUY Havells 100 @ 1501.25 on 2026-07-15 note=top-up",
                "SELL RELIANCE 50 @ 1300 on 2026-07-20",
                "UPDATE SECURITY HAVELLS sector=Consumer industry=Electrical",
                "SELL_SINCE stock=ExampleCo sell_date=2026-01-15 sell_price=120 note=manual",
            ]
        )
    )
    assert result.errors == []
    assert len(result.edits) == 4
    buy = result.edits[0]
    assert buy.action == "append_transaction"
    assert buy.fields["event_type"] == "Buy"
    assert buy.fields["quantity"] == 100
    assert buy.fields["price"] == "1501.25"
    sell = result.edits[1]
    assert sell.fields["event_type"] == "Sell"
    assert sell.fields["quantity"] == -50
    sec = result.edits[2]
    assert sec.action == "update_security"
    assert sec.fields["sector"] == "Consumer"
    ss = result.edits[3]
    assert ss.action == "append_sell_since"
    assert ss.fields["stock"] == "ExampleCo"


def test_parse_rejects_garbage() -> None:
    result = parse_prompt("please add some trades somehow")
    assert result.edits == []
    assert result.errors
    assert "Unrecognized" in result.errors[0]


def test_parse_kind_mismatch() -> None:
    result = parse_prompt(
        "BUY Havells 10 on 2026-07-15",
        kind=MasterKind.SECURITY,
    )
    assert result.edits == []
    assert any("targets transactions" in e for e in result.errors)


def _write_txn_fixture(path: Path) -> None:
    wb = openpyxl.Workbook()
    # Extra sheet first (like production V5)
    wb.active.title = "READ_ME"
    wb.active["A1"] = "notes"
    ws = wb.create_sheet("Sheet1")
    ws.append(
        [
            "Sr No",
            "Stock Name",
            "Date",
            "Buy/Sell",
            "Quantity",
            "Price",
            "Amount",
            "Source / Notes",
        ]
    )
    ws.append([1, "Havells", date(2024, 1, 2), "Buy", 10, 100, 1000, "seed"])
    wb.save(path)
    wb.close()


def _write_security_fixture(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Security Master"
    ws.append(
        [
            "Security ID",
            "Portfolio Name",
            "Canonical Company Name",
            "Current NSE Symbol",
            "Historical NSE Symbol",
            "BSE Code",
            "ISIN",
            "Current Status",
            "Corporate / Name History Notes",
            "Sector",
            "Industry",
            "Verification Status",
        ]
    )
    ws.append(
        [
            "SEC999",
            "Havells",
            "Havells India Limited",
            "HAVELLS",
            None,
            "500000",
            "INE000000000",
            "Active",
            None,
            "OldSector",
            "OldIndustry",
            "Verified",
        ]
    )
    wb.save(path)
    wb.close()


def _write_sell_since_fixture(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(
        [
            "Stock",
            "Sell Date",
            "Sell Price",
            "Then Portfolio Value",
            None,
            "Current Date",
            "Current Price",
            "Current Portfolio Value",
            None,
            "Stock Performance",
            "Portfolio Performance",
            "Outperformace",
            "Source / Corporate Action Note",
        ]
    )
    wb.save(path)
    wb.close()


@pytest.fixture
def master_dir(tmp_path: Path) -> Path:
    _write_security_fixture(tmp_path / "SECURITY_MASTER_V1.xlsx")
    _write_txn_fixture(tmp_path / "TRANSACTIONS_MASTER_V5_HARD_RECON_FIXES.xlsx")
    _write_sell_since_fixture(tmp_path / "Sell_Since_Completed_All_Missing_Stocks.xlsx")
    return tmp_path


def test_preview_transactions(master_dir: Path) -> None:
    preview = preview_workbook(MasterKind.TRANSACTIONS, master_dir=master_dir, limit=10)
    assert preview.sheet == "Sheet1"
    assert "Stock Name" in preview.columns
    assert preview.total_rows == 1
    assert preview.rows[0]["Stock Name"] == "Havells"


def test_apply_appends_transaction_and_updates_security(session, master_dir: Path, monkeypatch) -> None:
    raw = master_dir / "raw"
    (raw / "security_master").mkdir(parents=True)
    (raw / "transactions").mkdir(parents=True)
    monkeypatch.setattr(
        "pms_platform.config.settings.raw_data_dir",
        raw,
    )

    parsed = parse_prompt(
        "\n".join(
            [
                "BUY Havells 5 @ 200 on 2026-07-16 note=test",
                "UPDATE SECURITY HAVELLS sector=Consumer",
                "SELL_SINCE stock=Demo sell_date=2026-02-01 sell_price=99 note=x",
            ]
        )
    )
    assert not parsed.errors

    result = apply_edits(
        session,
        parsed.edits,
        master_dir=master_dir,
        reimport=False,
    )
    assert result.applied == 3
    assert result.errors == []
    assert len(result.backups) == 3

    txn_path = resolve_master_path(MasterKind.TRANSACTIONS, master_dir)
    wb = openpyxl.load_workbook(txn_path, data_only=True)
    ws = wb["Sheet1"]
    assert ws.cell(3, 2).value == "Havells"
    assert ws.cell(3, 4).value == "Buy"
    assert ws.cell(3, 5).value == 5
    wb.close()

    sec_path = resolve_master_path(MasterKind.SECURITY, master_dir)
    wb = openpyxl.load_workbook(sec_path, data_only=True)
    ws = wb["Security Master"]
    assert ws.cell(2, 10).value == "Consumer"
    wb.close()

    ss_path = resolve_master_path(MasterKind.SELL_SINCE, master_dir)
    wb = openpyxl.load_workbook(ss_path, data_only=True)
    ws = wb["Sheet1"]
    assert ws.cell(2, 1).value == "Demo"
    assert float(ws.cell(2, 3).value) == 99.0
    wb.close()

    assert (raw / "transactions" / "MASTER_TRANSACTIONS_V1.xlsx").is_file()
    assert (raw / "security_master" / "SECURITY_MASTER_V1.xlsx").is_file()
