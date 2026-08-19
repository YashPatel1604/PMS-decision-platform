"""Tests for Research Fair Value watchlist seed."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from pms_platform.watchlists import seed as seed_wl
from pms_platform.watchlists import service as wl


def _write_seed_xlsx(path: Path) -> Path:
    workbook = Workbook()
    mar16 = workbook.active
    mar16.title = "Mar16"
    mar16.append(["Company", "Equity"])
    mar16.append([None, None])
    mar16.append(["TestCo", 10])
    mar16.append(["Ambika Cotton", 5])
    fair = workbook.create_sheet("FairValue")
    fair.append([None, datetime(2016, 3, 1)])
    fair.append(["TestCo"])
    fair.append(["Competitors"])
    fair.append(["Previous Owned"])
    fair.append(["Heritage"])
    jun = workbook.create_sheet("Jun25")
    jun.append(["Company"])
    jun.append(["Supriya"])
    workbook.save(path)
    workbook.close()
    return path


def test_unique_company_names_skips_junk_and_unions_sheets(tmp_path: Path) -> None:
    path = _write_seed_xlsx(tmp_path / "watchlist.xlsx")
    names = seed_wl.unique_company_names(path)
    assert names == ["TestCo", "Ambika Cotton", "Heritage", "Supriya"]


def test_seed_watchlist_creates_default_and_refuses_overwrite(
    session, sample_security, tmp_path: Path
) -> None:
    path = _write_seed_xlsx(tmp_path / "watchlist.xlsx")
    with (
        patch("pms_platform.watchlists.resolution.resolve_bse_code", return_value=None),
        patch("pms_platform.watchlists.resolution.YahooFinanceClient") as yahoo,
    ):
        yahoo.return_value.search.return_value = []
        result = seed_wl.seed_watchlist(session, path=path)
        session.commit()

        assert result.watchlist_name == "Fair Value"
        assert result.added == 4
        assert result.skipped == 0
        members = wl.list_members(session, result.watchlist_id)
        assert {m.display_name for m in members} == {
            "TestCo",
            "Ambika Cotton",
            "Heritage",
            "Supriya",
        }
        testco = next(m for m in members if m.display_name == "TestCo")
        assert testco.security_id == "SEC999"

        with pytest.raises(wl.WatchlistError, match="already has"):
            seed_wl.seed_watchlist(session, path=path)

        again = seed_wl.seed_watchlist(session, path=path, force=True)
        assert again.added == 0
        assert again.skipped == 4
